"""Weekly Pulse review workbook (v2, 2026-10-27, after Darryl's width review).

Generates Reviews/YYYY-MM-DD_Pulse_Review.xlsx and a manifest JSON of the values as sent,
which the Monday read-back diffs against. One column per field, pre-filled: a correct row
costs nothing. Edit columns sit next to the issue; reference columns on the right.
The workbook is a weekly form, never a register.
"""
import sys, json, datetime as dt
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import eos_registers as R, eos_pulse as P
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

ROOT = Path(R.ROOT)
REVIEWERS = json.loads((ROOT / "Registers" / "pulse" / "_reviewers.json").read_text())
VERDICTS = ["No change", "Resolved", "Not an issue", "Needs Darryl"]
SEVERITIES = ["High", "Medium", "Low"]
SEV_RANK = {"High": 0, "Medium": 1, "Low": 2}

def reviewers_for(track, product):
    t = REVIEWERS.get(track, {})
    return t.get(product) or t.get("*") or []
ALL_REVIEWERS = sorted({n for t in ("development", "delivery") for v in REVIEWERS[t].values() for n in v})

# column spec: (header, width, editable)
COLS = [("ID", 6, False), ("Issue", 52, False), ("Owner", 20, True), ("Sev", 8, True), ("Next action", 40, True),
        ("Verdict", 13, True), ("Comment", 26, True), ("Reviewed by", 18, True),
        ("Team", 7, False), ("Product", 9, False), ("Account", 12, False), ("Heard", 10, False)]

def rows_from_pulse():
    forums = {f["key"]: f for f in P.load_forums()}
    pulse = [m for m, body, p in R.load_items("pulse")]
    rows, raw = [], set()
    for m in pulse:
        raw.update(p.strip() for p in (m.get("owner", "") or "").replace(";", ",").split(",") if p.strip())
        if m.get("status") != "live" or P.is_knowledge(m): continue
        f = forums.get(m.get("forum"), {})
        rows.append(dict(id=m["id"], title=m.get("title", ""), owner=m.get("owner", "") or "", sev=m.get("severity", ""),
                         nxt=m.get("next_action", "") or "", team=f.get("track") or "unassigned", product=m.get("product", ""),
                         account=m.get("account", ""), heard=m.get("last_seen", ""), forum=f.get("name", "")))
    rows.sort(key=lambda r: (r["team"], SEV_RANK.get(r["sev"], 3), r["id"]))
    people = set(ALL_REVIEWERS)
    for p in raw:
        fx = REVIEWERS["people_fixes"].get(p, p); people.update(fx if isinstance(fx, list) else [fx])
    return rows, sorted(people)

def build(out=None, today=None):
    today = today or dt.date.today().isoformat()
    rows, people = rows_from_pulse()
    wb = Workbook(); ws = wb.active; ws.title = "Review"
    head = PatternFill("solid", fgColor="1F3864"); grey = PatternFill("solid", fgColor="F2F2F2"); blue = PatternFill("solid", fgColor="E8F1FB")
    ws.append([c[0] for c in COLS])
    for i, (h, w, e) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = head; c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r["id"], r["title"], r["owner"], r["sev"], r["nxt"], "No change", "", "", r["team"][:3].title(),
                   r["product"], r["account"], r["heard"]])
        rv = " or ".join(reviewers_for(r["team"], r["product"]))
        ws.cell(ws.max_row, 1).comment = Comment(f"Reviewers: {rv}\nRaised in: {r['forum']}", "EOS")
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            e = COLS[c.column - 1][2]
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.fill = blue if e else grey
            c.protection = Protection(locked=not e)
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n}"
    # Orig sheet (hidden) holds the sent values; a changed row is highlighted in the sheet itself.
    og = wb.create_sheet("Orig"); og.sheet_state = "hidden"
    og.append(["ID", "Owner", "Sev", "Next action", "Verdict"])
    for r in rows: og.append([r["id"], r["owner"], r["sev"], r["nxt"], "No change"])
    changed = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(COLS))}{n}",
        FormulaRule(formula=[f'OR($C2<>Orig!$B2,$D2<>Orig!$C2,$E2<>Orig!$D2,$F2<>Orig!$E2,LEN($G2)>0)'], fill=changed))
    ls = wb.create_sheet("Lists"); ls.sheet_state = "hidden"
    for i, v in enumerate(VERDICTS, 1): ls.cell(i, 1, v)
    for i, v in enumerate(SEVERITIES, 1): ls.cell(i, 2, v)
    for i, v in enumerate(people, 1): ls.cell(i, 3, v)
    for i, v in enumerate(ALL_REVIEWERS, 1): ls.cell(i, 4, v)
    def dv(sheet, col, ref, cnt, last):
        d = DataValidation(type="list", formula1=f"=Lists!${ref}$1:${ref}${cnt}", allow_blank=True, showErrorMessage=True)
        sheet.add_data_validation(d); d.add(f"{col}2:{col}{last}")
    dv(ws, "F", "A", len(VERDICTS), n); dv(ws, "C", "C", len(people), n); dv(ws, "D", "B", len(SEVERITIES), n); dv(ws, "H", "D", len(ALL_REVIEWERS), n)
    ws.protection.sheet = True; ws.protection.autoFilter = False; ws.protection.sort = False; ws.protection.formatRows = False

    nw = wb.create_sheet("New issues")
    nw.append(["Team", "Issue (one line)", "Owner", "Sev", "Raised by"])
    for i, w in enumerate([10, 70, 20, 8, 18], 1):
        nw.column_dimensions[get_column_letter(i)].width = w
        c = nw.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = head
    d = DataValidation(type="list", formula1='"Delivery,Development"', allow_blank=True); nw.add_data_validation(d); d.add("A2:A50")
    dv(nw, "C", "C", len(people), 50); dv(nw, "D", "B", len(SEVERITIES), 50); dv(nw, "E", "D", len(ALL_REVIEWERS), 50)

    hu = wb.create_sheet("How to use", 0); hu.column_dimensions["A"].width = 90
    lines = [f"Pulse Review, week of {today}. Return by Friday.", "",
        "1. On the Review tab, filter Team (and Product if you want) to your rows. Grey is the record; blue is yours.",
        "2. If a row is right, do nothing. Every blue cell is already filled with what the record says.",
        "3. If a row is wrong, change the blue cell that is wrong: Owner, Sev, or Next action. The row turns yellow so you can see what you changed.",
        "4. Verdict: Resolved (it is done), Not an issue (it should not be tracked), Needs Darryl (say why in Comment). Otherwise leave No change.",
        "5. Put your name in Reviewed by on every row you changed. Hover the ID to see who reviews that row and where it was raised.",
        "6. Anything we missed goes on the New issues tab, one line each.", "",
        "Ground rules: work in the shared copy only, never a download or an email copy; AutoSave on; do not add columns or tabs, sort, or delete rows; "
        "one owner per row; back by Friday close of business. Changes are filed on Monday against each issue with your name and the date."]
    for i, t in enumerate(lines, 1):
        c = hu.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1: c.font = Font(bold=True, size=13)
    out = Path(out) if out else ROOT / "Reviews" / f"{today}_Pulse_Review.xlsx"
    out.parent.mkdir(exist_ok=True); wb.save(out)
    manifest = {"date": today, "rows": {r["id"]: {"owner": r["owner"], "sev": r["sev"], "nxt": r["nxt"], "team": r["team"], "product": r["product"]} for r in rows}}
    out.with_suffix(".manifest.json").write_text(json.dumps(manifest, indent=1))
    teams = {}
    for r in rows: teams[r["team"]] = teams.get(r["team"], 0) + 1
    return out, teams, len(people)

if __name__ == "__main__":
    out, teams, np = build(today=sys.argv[1] if len(sys.argv) > 1 else None)
    print(out); print(teams, "people:", np, "width:", sum(c[1] for c in COLS))


# ----------------------------------------------------------------------------- read-back (2026-10-27)

from openpyxl import load_workbook

RETURNED = ROOT / "Reviews" / "returned"
LOG = ROOT / "Reviews" / "_log.json"
REVIEW_FORUM = {"delivery": "review-delivery", "development": "review-development"}

def _log():
    try: return json.loads(LOG.read_text())
    except Exception: return {"sent": [], "returned": []}

def _save_log(d): LOG.write_text(json.dumps(d, indent=1))

def note_sent(date):
    d = _log()
    if date not in [s["date"] for s in d["sent"]]:
        d["sent"].append({"date": date, "at": dt.datetime.now().isoformat(timespec="minutes")}); _save_log(d)

def latest_manifest():
    ms = sorted((ROOT / "Reviews").glob("*_Pulse_Review.manifest.json"))
    return ms[-1] if ms else None

def read_back(xlsx_path) -> dict:
    """Diff a returned workbook against the manifest it was generated from. Pure read; nothing is written."""
    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["Review"]
    hdr = [c.value for c in ws[1]]
    want = ["ID", "Issue", "Owner", "Sev", "Next action", "Verdict", "Comment", "Reviewed by", "Team"]
    if hdr[:9] != want:
        return {"error": f"Unexpected columns: {hdr[:9]}. The sheet structure was changed; expected {want}."}
    # manifest: match the file's date prefix, else the latest
    mf = ROOT / "Reviews" / (xlsx_path.name[:10] + "_Pulse_Review.manifest.json")
    if not mf.exists(): mf = latest_manifest()
    man = json.loads(mf.read_text()) if mf else {"date": "", "rows": {}}
    changes, rejected, seen = [], [], set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        rid = (r[0] or "").strip() if isinstance(r[0], str) else r[0]
        if not rid: continue
        seen.add(rid)
        base = man["rows"].get(rid)
        if base is None:
            rejected.append({"id": rid, "reason": "not in the workbook as sent"}); continue
        owner, sev, nxt = (r[2] or "").strip(), (r[3] or "").strip(), (r[4] or "").strip()
        verdict, comment, who = (r[5] or "No change").strip(), (r[6] or "").strip(), (r[7] or "").strip()
        diff = {}
        if owner != base["owner"]: diff["owner"] = [base["owner"], owner]
        if sev != base["sev"]: diff["severity"] = [base["sev"], sev]
        if nxt != base["nxt"]: diff["next_action"] = [base["nxt"], nxt]
        if verdict != "No change": diff["verdict"] = ["No change", verdict]
        if comment: diff["comment"] = ["", comment]
        if not diff: continue
        allowed = reviewers_for(base["team"], base["product"])
        c = {"id": rid, "title": r[1], "team": base["team"], "product": base["product"], "reviewer": who, "diff": diff,
             "warnings": []}
        if not who: c["warnings"].append("no reviewer name; pick one or skip")
        elif who not in ALL_REVIEWERS: c["warnings"].append(f"{who} is not on the reviewer list")
        elif who not in allowed: c["warnings"].append(f"{who} is not a reviewer for {base['product']} ({' or '.join(allowed)})")
        if verdict == "Needs Darryl" and not comment: c["warnings"].append("Needs Darryl without a comment")
        changes.append(c)
    new = []
    if "New issues" in wb.sheetnames:
        for r in wb["New issues"].iter_rows(min_row=2, values_only=True):
            if not (r[1] or "").strip(): continue
            new.append({"team": (r[0] or "").strip().lower(), "title": (r[1] or "").strip(), "owner": (r[2] or "").strip(),
                        "sev": (r[3] or "Medium").strip(), "by": (r[4] or "").strip(),
                        "warnings": ([] if (r[4] or "").strip() in ALL_REVIEWERS else ["no valid Raised by name"]) +
                                    ([] if (r[0] or "").strip().lower() in REVIEW_FORUM else ["team not set"])})
    missing = sorted(set(man["rows"]) - seen)
    by = {}
    for c in changes: by[c["reviewer"] or "(unattributed)"] = by.get(c["reviewer"] or "(unattributed)", 0) + 1
    return {"file": xlsx_path.name, "manifest_date": man["date"], "rows_sent": len(man["rows"]), "rows_seen": len(seen),
            "missing_rows": missing, "changes": changes, "new": new, "rejected": rejected, "by_reviewer": by}

def apply(preview: dict, accept: list[dict], today=None) -> dict:
    """accept: [{"id": "P030", "reviewer": "Ron Snyder"}] for changes; [{"new": 0, "reviewer": ...}] for new issues.
    Writes dated, attributed lines and field changes; caller regenerates and commits."""
    today = today or dt.date.today().isoformat()
    acc = {a["id"]: a for a in accept if "id" in a}
    accn = {a["new"]: a for a in accept if "new" in a}
    done, created = [], []
    for c in preview["changes"]:
        a = acc.get(c["id"])
        if not a: continue
        who = (a.get("reviewer") or c["reviewer"] or "").strip()
        if not who: continue
        f = R.find_item(c["id"])
        if not f: continue
        meta, body, _p = f
        d = c["diff"]; parts = []
        if "owner" in d: meta["owner"] = d["owner"][1]; parts.append(f"owner {d['owner'][0] or 'none'} -> {d['owner'][1] or 'none'}")
        if "severity" in d: meta["severity"] = d["severity"][1]; parts.append(f"severity {d['severity'][0]} -> {d['severity'][1]}")
        if "next_action" in d: meta["next_action"] = d["next_action"][1]; parts.append(f"next action: {d['next_action'][1]}")
        v = d.get("verdict", [None, None])[1]
        if v == "Resolved": meta["status"] = "resolved"; parts.append("RESOLVED")
        elif v == "Not an issue": meta["status"] = "resolved"; meta["disposition_note"] = f"Not an issue per {who} (weekly review {today})"; parts.append("NOT AN ISSUE")
        elif v == "Needs Darryl":
            if meta.get("escalation") not in ("escalated",): meta["escalation"] = "candidate"; meta["escalation_test"] = "asked"
            parts.append("NEEDS DARRYL")
        if "comment" in d: parts.append(f"comment: {d['comment'][1]}")
        meta["last_touched"] = today
        body = body.rstrip("\n") + f"\n\n{today} ({who}, review): " + "; ".join(parts) + ".\n"
        R.write_item(meta, body); done.append({"id": c["id"], "reviewer": who, "what": "; ".join(parts)})
    for i, n in enumerate(preview.get("new", [])):
        a = accn.get(i)
        if not a: continue
        who = (a.get("reviewer") or n["by"] or "").strip()
        if not who or n["team"] not in REVIEW_FORUM: continue
        pid = R.next_id("P")
        meta = {"id": pid, "kind": "pulse", "status": "live", "title": n["title"], "owner": n["owner"], "product": a.get("product") or "entroCIM",
                "severity": n["sev"] or "Medium", "forum": REVIEW_FORUM[n["team"]], "next_action": "", "created": today, "last_touched": today,
                "needs_triage": False, "account": "", "customer_visible": False, "escalated_to": "", "escalation": "", "escalation_test": "",
                "first_seen": today, "last_seen": today, "links": [], "people": [p.strip() for p in n["owner"].replace(";", ",").split(",") if p.strip()],
                "sightings": 1, "source": "weekly review"}
        body = f"{today} ({who}, review): Raised on the New issues tab of the weekly Pulse review: {n['title']}\n"
        R.write_item(meta, body); created.append({"id": pid, "title": n["title"], "reviewer": who})
    d = _log(); d["returned"].append({"date": today, "file": preview.get("file"), "manifest_date": preview.get("manifest_date"),
                                     "applied": len(done), "created": len(created), "by_reviewer": preview.get("by_reviewer", {})}); _save_log(d)
    return {"applied": done, "created": created}
