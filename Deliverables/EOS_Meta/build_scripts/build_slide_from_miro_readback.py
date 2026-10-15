# Build one native PowerPoint slide on the Wesco template from a Miro connector read-back (Canvas Composer SVG).
# Usage: python3 build_slide_from_miro_readback.py readback.svg template.pptx out.pptx "<Slide title>" <region_rect_id>
# The region rect (the 16:9 option card on the board) defines what maps onto the slide content area.
# Its own title/subtitle text areas are dropped: the slide title placeholder carries the title, and subtitles are not used.
# Rects become rounded rectangles (fill, stroke, text from data-content), connectors become single native shapes:
# straight side-to-side links as straight connectors, bottom-to-bottom links as one elbow freeform. Positions scale uniformly.
import sys, re, html
from lxml import etree
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

svg_path, tpl, out, slide_title, region_id = sys.argv[1:6]
FONT='Arial'
X0, Y0, WMAX, YMAX = 0.73, 1.75, 11.86, 6.55          # slide content area, inches

def rgb(h):
    h=h.lstrip('#')[:6]; return RGBColor.from_string(h.upper())
def is_none(c): return c is None or c.lower() in ('none','#ffffffff','transparent')

root=etree.parse(svg_path).getroot()
def tag(e): return etree.QName(e).localname
# collect rects with absolute coords (apply group translate)
items={}
def walk(e,dx=0,dy=0):
    t=e.get('transform')
    if t:
        m=re.match(r'translate\(([-\d.]+),([-\d.]+)\)',t); dx+=float(m.group(1)); dy+=float(m.group(2))
    if tag(e) in ('rect','textArea') and e.get('data-type')!='frame' and e.get('id'):
        items[e.get('id')]=dict(kind=tag(e),x=float(e.get('x'))+dx,y=float(e.get('y'))+dy,w=float(e.get('width')),h=float(e.get('height')),el=e)
    for c in e: walk(c,dx,dy)
walk(root)
links=[e for e in root.iter() if tag(e)=='line' and e.get('data-start')]

reg=items[region_id]
inside=lambda it: it['x']>=reg['x']-1 and it['y']>=reg['y']-1 and it['x']+it['w']<=reg['x']+reg['w']+1 and it['y']+it['h']<=reg['y']+reg['h']+1
shapes={k:v for k,v in items.items() if k!=region_id and inside(v) and v['kind']=='rect'}
# content bounds = shapes only (text areas in the region are the board title/subtitle, carried by the slide title)
bx0=min(v['x'] for v in shapes.values()); by0=min(v['y'] for v in shapes.values())
bx1=max(v['x']+v['w'] for v in shapes.values()); by1=max(v['y']+v['h'] for v in shapes.values())
S=min(WMAX/(bx1-bx0),(YMAX-Y0)/(by1-by0))              # inches per board px
PT=lambda px: round(px*S*72*2)/2                         # board px font -> pt, to 0.5
mx=lambda x: X0+(x-bx0)*S; my=lambda y: Y0+(y-by0)*S

prs=Presentation(tpl)
ids=prs.slides._sldIdLst
for sid in list(ids):
    prs.part.drop_rel(sid.get(qn('r:id'))); ids.remove(sid)
lay={l.name:l for l in prs.slide_layouts}['Blank-with-Title']
s=prs.slides.add_slide(lay); s.shapes.title.text=slide_title

geo={}
for k,v in shapes.items():
    e=v['el']; x,y,w,h=mx(v['x']),my(v['y']),v['w']*S,v['h']*S; geo[k]=(x,y,w,h)
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,In(x),In(y),In(w),In(h))
    rx=float(e.get('rx') or 0); sh.adjustments[0]=min(0.5,rx/min(v['w'],v['h'])) if rx else 0
    if is_none(e.get('fill')): sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(e.get('fill'))
    if is_none(e.get('stroke')): sh.line.fill.background()
    else: sh.line.color.rgb=rgb(e.get('stroke')); sh.line.width=Pt(1.25)
    sh.shadow.inherit=False
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for a in ('margin_left','margin_right'): setattr(tf,a,In(0.08))
    tf.margin_top=tf.margin_bottom=In(0.03)
    content=html.unescape(e.get('data-content') or '')
    lines=re.split(r'<br\s*/?>',content)
    size=PT(float(e.get('data-font-size') or 15)); color=rgb(e.get('data-text-color') or '#1d252d')
    allbold=e.get('data-font-weight')=='bold'
    for i,l in enumerate(lines):
        bold=allbold or '<b>' in l; t=re.sub('<[^>]+>','',l).strip()
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.CENTER; p.space_after=Pt(1)
        r=p.add_run(); r.text=t; r.font.name=FONT; r.font.color.rgb=color; r.font.bold=bold
        r.font.size=Pt(size+2.5 if (bold and len(lines)>1 and i==0) else size)   # card headings one step up

def side(k,sd):
    x,y,w,h=geo[k]
    return {'right':(x+w,y+h/2),'left':(x,y+h/2),'bottom':(x+w/2,y+h),'top':(x+w/2,y)}[sd]
def arrowhead(ln): ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
STAND=0.07
connectors=[]
for e in links:
    a,b=e.get('data-start'),e.get('data-end')
    if a not in geo or b not in geo: continue
    (x1,y1),(x2,y2)=side(a,e.get('data-start-side')),side(b,e.get('data-end-side'))
    color=rgb(e.get('stroke') or '#313131'); dashed=bool(e.get('stroke-dasharray')); label=e.get('data-content')
    if e.get('data-start-side')=='bottom' and e.get('data-end-side')=='bottom':
        ydrop=max(y1,y2)+0.30
        ff=s.shapes.build_freeform(In(x1),In(y1),scale=1.0)
        ff.add_line_segments([(In(x1),In(ydrop)),(In(x2),In(ydrop)),(In(x2),In(y2+STAND))],close=False)
        c=ff.convert_to_shape(); c.fill.background()
        if label:
            tb=s.shapes.add_textbox(In(min(x1,x2)),In(ydrop+0.03),In(abs(x1-x2)),In(0.22))
            pts=[(x1,y1),(x1,ydrop),(x2,ydrop),(x2,y2+STAND)]
    else:
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(x1),In(y1),In(x2-STAND),In(y2))
        pts=[(x1,y1),(x2-STAND,y2)]
        if label:
            gap=x2-x1; tb=s.shapes.add_textbox(In(x1),In(y1-0.24),In(gap),In(0.2))
    c.line.color.rgb=color; c.line.width=Pt(1.5)
    if dashed: c.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    arrowhead(c.line._get_or_add_ln()); c.shadow.inherit=False
    connectors.append((a,b,pts))
    if label:
        tf=tb.text_frame; tf.word_wrap=True
        for at in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,at,0)
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=label
        r.font.size=Pt(9); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=color

s.notes_slide.notes_text_frame.text=('Test build from the Miro connector read-back (board "Launch Studio Sean Deck 2026-10", option 1A), 2026-10-15. '
    'Positions, sizes, colors, text and connectors come from the board, scaled onto the slide. Not for distribution: the board wording predates the 2026-10-14 deck revision.')
prs.save(out)

# Deck standards §2: every connector endpoint must land on a node edge (arrowhead standoff 0.07 in allowed)
def on_edge(px,py,tol=0.09):
    for k,(x,y,w,h) in geo.items():
        inx=x-tol<=px<=x+w+tol; iny=y-tol<=py<=y+h+tol
        if inx and iny and (min(abs(px-x),abs(px-x-w))<=tol or min(abs(py-y),abs(py-y-h))<=tol): return k
    return None
ok=True
for a,b,pts in connectors:
    s0,e0=on_edge(*pts[0]),on_edge(*pts[-1])
    print(f'{a[-4:]}->{b[-4:]}: start on {s0 and s0[-4:]}, end on {e0 and e0[-4:]}')
    ok&= s0==a and e0==b
print('scale in/px',round(S,5),'content bottom',round(my(by1),2),'ENDPOINTS OK' if ok else 'ENDPOINT DEFECT')
