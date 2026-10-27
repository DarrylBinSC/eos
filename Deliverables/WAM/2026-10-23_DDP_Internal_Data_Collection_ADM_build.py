# DDP | Internal Data Collection, native rebuild of the 2026-10-19 markup for Alan Martin.
# Run from a folder containing template.pptx (Wesco template): python3 build_ddp_slide.py -> 2026-10-23_DDP_Internal_Data_Collection_ADM.pptx
# Red = proposed additions from Darryl Benson's markup (CONSUME column, Markup notes, decision numbers 8 and 9).

from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

INK=C(0x2B,0x2B,0x2B); GREY=C(0x59,0x59,0x59); LINE=C(0x3F,0x3F,0x3F)
RED=C(0xC0,0x00,0x1E); WHITE=C(0xFF,0xFF,0xFF); BLACK=C(0x00,0x00,0x00)
F='Arial'

prs=Presentation('template.pptx')
ids=prs.slides._sldIdLst
for sid in list(ids):
    rId=sid.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId); ids.remove(sid)
L={l.name:l for l in prs.slide_layouts}
s=prs.slides.add_slide(L['Blank-with-Title'])
s.shapes.title.text_frame.paragraphs[0].add_run().text='DDP | Internal Data Collection'

def frame(tf,margin=0.06):
    tf.word_wrap=True
    tf.margin_left=tf.margin_right=In(margin); tf.margin_top=tf.margin_bottom=In(0.04)
def para(tf,text,size,bold=False,color=INK,first=False,align=PP_ALIGN.CENTER,bullet=False,space=1.5,runs=None):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space)
    if bullet:
        pPr=p._p.get_or_add_pPr(); pPr.set('marL','114300'); pPr.set('indent','-114300')
        for tag,attrs in (('a:buClr',None),('a:buSzPct',{'val':'100000'}),('a:buFont',{'typeface':'Arial'}),('a:buChar',{'char':'•'})):
            el=pPr.makeelement(qn(tag),attrs or {})
            if tag=='a:buClr':
                c=el.makeelement(qn('a:srgbClr'),{'val':'%02X%02X%02X'%tuple(color)}); el.append(c)
            pPr.append(el)
    for t,b,c in (runs or [(text,bold,color)]):
        r=p.add_run(); r.text=t; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=c; r.font.name=F
    return p
def box(x,y,w,h,lines,size=8,line=LINE,lw=1.0,fill=WHITE,rounded=True,bullets=False,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,name=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,In(x),In(y),In(w),In(h))
    if rounded: sh.adjustments[0]=0.06
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=line; sh.line.width=Pt(lw); sh.shadow.inherit=False
    tf=sh.text_frame; frame(tf); tf.vertical_anchor=anchor
    for i,ln in enumerate(lines):
        if isinstance(ln,tuple): t,b,c=ln
        else: t,b,c=ln,False,INK
        para(tf,t,size,b,c,first=(i==0),align=(PP_ALIGN.LEFT if bullets else align),bullet=bullets,space=(0.5 if bullets and size<7 else 1.5))
    if name: sh.name=name
    return sh
def label(x,y,w,h,text,size=9,color=GREY,bold=False,align=PP_ALIGN.CENTER):
    tb=s.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=tb.text_frame; frame(tf,0)
    tf.vertical_anchor=MSO_ANCHOR.BOTTOM
    para(tf,text,size,bold,color,first=True,align=align); return tb
def arrow(x1,y1,x2,y2,color=LINE,lw=1.0):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(x1),In(y1),In(x2),In(y2))
    c.line.color.rgb=color; c.line.width=Pt(lw)
    ln=c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    return c
def path(pts,color=LINE,lw=1.0):
    ff=s.shapes.build_freeform(In(pts[0][0]),In(pts[0][1]),scale=1.0)
    ff.add_line_segments([(In(x),In(y)) for x,y in pts[1:]],close=False)
    sh=ff.convert_to_shape(); sh.fill.background(); sh.line.color.rgb=color; sh.line.width=Pt(lw); sh.shadow.inherit=False
    ln=sh.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
    return sh
def lane(x,y,w,h,text):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,In(x),In(y),In(w),In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=LINE; sh.line.width=Pt(1.0); sh.shadow.inherit=False
    tf=sh.text_frame; frame(tf,0.02); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf._txBody.bodyPr.set('vert','vert270')
    para(tf,text,8,False,INK,first=True); return sh
def circle(cx,cy,d,text):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,In(cx-d/2),In(cy-d/2),In(d),In(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=BLACK; sh.line.fill.background(); sh.shadow.inherit=False
    tf=sh.text_frame; frame(tf,0); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    para(tf,text,8,False,WHITE,first=True); return sh

# ---- geometry (inches) ----
G=0.12
X_WHAT,X_WHEN,X_HOW,X_WHERE=1.65,3.02,4.39,5.76; W=1.25
X_EUDM,X_ADM=7.15,8.35; WC=1.08
X_END=9.75; X_PANEL=10.19; WP=2.40
Y_HDR=1.42
# lanes
S0,S1=1.70,3.05      # Serialization
A0,A1=3.20,5.62      # Additional attributes
V0,V1=5.74,6.44      # VAS

# column headers
for x,t in ((X_WHAT,'WHAT'),(X_WHEN,'WHEN'),(X_HOW,'HOW'),(X_WHERE,'WHERE')):
    label(x,Y_HDR,W,0.24,t,9.5)
hdr=label(X_EUDM,Y_HDR-0.12,X_ADM+WC-X_EUDM,0.36,'CONSUME',9.5,RED,True)
para(hdr.text_frame,'(proposed addition)',6.5,False,RED,align=PP_ALIGN.CENTER,space=0)
label(X_PANEL,Y_HDR,WP,0.24,'KEY PROCESS DESIGN DECISIONS',8.5)

# lanes
lane(1.30,S0,0.28,S1-S0,'Serialization')
lane(1.30,A0,0.28,A1-A0,'Additional Data Collection Attributes')
lane(1.30,V0,0.28,V1-V0,'VAS')

# Serialization row
b_what1=box(X_WHAT,1.70,W,0.64,['Part number is flagged for Serialization P360'])
b_what2=box(X_WHAT,2.42,W,0.63,['Currently, this is managed in P360 and is not warehouse-specific'])
b_when1=box(X_WHEN,1.70,W,0.64,['Currently, this can be done at inbound receipt, pick, work order completion, or pack'],size=7)
b_how1=box(X_HOW,1.70,W,0.62,['Serial capture is performed manually or via scan in WMX'])
b_how2=box(X_HOW,2.42,W,0.63,['Currently, manual capture occurs when components of an FG require serial capture'],size=7.5)
b_where1=box(X_WHERE,1.70,W,0.62,['Serial number capture is stored in WMX'])

# Additional attributes row
b_what3=box(X_WHAT,3.20,W,0.68,['If a part is flagged for serialization, then additional attributes should be able to be collected'],size=7)
b_what4=box(X_WHAT,3.96,W,1.08,['MAC Address','IP Address','Asset Tag','Firmware Version','Software Version','Time Zone','Pixel Test','DOA Check'],size=6.5,bullets=True)
b_what5=box(X_WHAT,5.10,W,0.50,['Should additional data collection attribute requirements be mastered?'],size=7)
b_when2=box(X_WHEN,3.20,W,0.70,['Should additional data collection occur in WMX at time of pick, work order completion, or pack?'],size=7.5)
b_how3=box(X_HOW,3.20,W,2.40,['MAC Address is manually entered when programming a device.','Asset tags are manually entered if there is no bar code','IP Addresses are customer-provided','All other attributes are manually entered'],size=6.5,bullets=True)
b_where2=box(X_WHERE,3.20,W,1.70,['Serial numbers, MAC Address and Asset Tag data should be stored in WMX and may be entered in a customer-provided document','All other attributes are stored on a customer-provided document'],size=6.5,bullets=True)

# VAS row
b_what6=box(X_WHAT,V0,W,V1-V0,['Should additional data collection requirements be defined as part of a VAS WI?'],size=7.5)
b_when3=box(X_WHEN,V0,W,V1-V0,['Should VAS integrate w/ WMX to create Tasks for each data collection requirement?'],size=7.5)
b_where3=box(X_WHERE,V0,W,V1-V0,['Should WMX integrate w/ VAS to return the data that was collected?'],size=7.5)

# Consume (proposed)
CY0,CY1=3.20,4.90
box(X_EUDM,CY0,WC,CY1-CY0,[('EUDM',True,RED),'WMX synchronizes the collected attributes to EUDM.'],size=7,line=RED,lw=1.75,anchor=MSO_ANCHOR.TOP)
box(X_ADM,CY0,WC,CY1-CY0,[('ADM',True,RED),'Consumes serial number, MAC address and asset tag from EUDM to build and maintain the asset record. ADM does not capture, and takes nothing directly from WMX or VAS.'],size=7,line=RED,lw=1.75,anchor=MSO_ANCHOR.TOP)

# Start / End
CYm=(CY0+CY1)/2
circle(0.955,2.01,0.42,'Start')
D_END=0.40
circle(X_END+D_END/2,CYm,D_END,'End')

# connectors (arrow tip stops ~0.05 short of target edge)
k=0.05
arrow(0.955+0.21,2.01,1.30-k,2.01)                                  # Start -> lane
arrow(X_WHAT+W,2.01,X_WHEN-k,2.01)                                  # what1 -> when1
arrow(X_WHEN+W,2.01,X_HOW-k,2.01)                                   # when1 -> how1
arrow(X_HOW+W,2.01,X_WHERE-k,2.01)                                  # how1 -> where1
arrow(X_WHAT+W/2,3.05,X_WHAT+W/2,3.20-k)                            # what2 -> what3
arrow(X_WHAT+W/2,3.88,X_WHAT+W/2,3.96-k)                            # what3 -> what4
arrow(X_WHAT+W/2,5.04,X_WHAT+W/2,5.10-k)                            # what4 -> what5
arrow(X_WHAT+W,3.54,X_WHEN-k,3.54)                                  # what3 -> when2
arrow(X_HOW+W/2,3.05,X_HOW+W/2,3.20-k)                              # how2 -> how3
arrow(X_WHAT+W,4.50,X_HOW-k,4.50)                                   # what4 -> how3
arrow(X_HOW+W,4.05,X_WHERE-k,4.05)                                  # how3 -> where2
arrow(X_WHERE+W/2,4.90,X_WHERE+W/2,V0-k)                            # where2 -> where3
arrow(X_WHERE+W,2.01,X_EUDM-k,CY0+0.18)                             # where1 -> EUDM (proposed)
arrow(X_WHERE+W,4.05,X_EUDM-k,4.05)                                 # where2 -> EUDM
arrow(X_EUDM+WC,CYm,X_ADM-k,CYm)                                    # EUDM -> ADM
arrow(X_ADM+WC,CYm,X_END-k,CYm)                                     # ADM -> End
path([(X_WHERE+W,(V0+V1)/2),(X_END+D_END/2,(V0+V1)/2),(X_END+D_END/2,CYm+D_END/2+k)])  # where3 -> End

# Decisions panel
pan=box(X_PANEL,1.70,WP,3.45,[],size=7,rounded=True,anchor=MSO_ANCHOR.TOP)
tf=pan.text_frame
dec=['Serialization is required if any data collection is required',
     'Should data collection requirements be mastered in P360 or be defined in VAS based upon a part + WI or Other?',
     'Should data collection be defined for a specific stage like serial numbers?',
     'Are there any data collection entry process requirements?',
     'If VAS is leveraged to define data collection, then should VAS create tasks in WMX for each data collection attribute?',
     'If VAS is leveraged to define data collection, then should WMX return the data collected?',
     'Consideration should be given to when a WI is required. We don’t want to force the use of WI when not necessary']
for i,t in enumerate(dec):
    para(tf,t,7,False,INK,first=(i==0),align=PP_ALIGN.LEFT,bullet=True,space=2.5)
para(tf,'',7,align=PP_ALIGN.LEFT,bullet=True,space=2.5,runs=[('8  ',True,RED),('Is the assumption that once WMX collects all data required, that WAM will consume the data from WMX or the DL?',False,INK)])
para(tf,'',7,align=PP_ALIGN.LEFT,bullet=True,space=2.5,runs=[('9  ',True,RED),('Are there any timing issues between when WMX would collect the data and when WAM or other apps may want to consume it?',False,INK)])
para(tf,'Other?',7,False,INK,align=PP_ALIGN.LEFT,bullet=True,space=0)
# clear the empty first paragraph left by box()
first=tf.paragraphs[0]
if not first.runs: first._p.getparent().remove(first._p)

# Markup notes (proposed)
mk=s.shapes.add_textbox(In(X_PANEL),In(5.28),In(WP),In(1.15)); mtf=mk.text_frame; frame(mtf,0.02)
para(mtf,'Markup notes',8,True,RED,first=True,align=PP_ALIGN.LEFT,space=2)
for t in ('Decision 8: ADM consumes from the DL.',
          'Decision 9: Timing issues are unknown at this point.',
          'Added design decision: which attributes must be stored in WMX, not only on a document, so they reach EUDM and ADM? Anything held only on a document never reaches ADM.'):
    para(mtf,t,6.5,False,RED,align=PP_ALIGN.LEFT,bullet=True,space=1.5)

# speaker note
s.notes_slide.notes_text_frame.text=('Native rebuild of Alan Martin’s DDP | Internal Data Collection slide with Darryl Benson’s 2026-10-19 markup. '
    'Red elements are the proposed additions: the CONSUME column (EUDM, ADM), decision numbers 8 and 9, and the Markup notes. '
    'Every box, lane and connector is a separate editable shape.')

prs.save('2026-10-23_DDP_Internal_Data_Collection_ADM.pptx')
print('saved')
