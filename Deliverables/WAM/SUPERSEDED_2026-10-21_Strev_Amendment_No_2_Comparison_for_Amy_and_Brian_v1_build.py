import re, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = "ref.docx"
OUT = "2026-10-21_Strev_Amendment_No_2_Comparison_for_Amy_and_Brian.docx"

CITE_COLOR = RGBColor(0x59, 0x59, 0x59)
EFFECT_FILL = "EAF1FB"   # pale blue, highlights the difference column
HEAD_FILL = "FFFFFF"
RULE_COLOR = "BFBFBF"
TABLE_PT = 9

doc = Document(TEMPLATE)

# ---- clear template body (keep the final sectPr) ----
body = doc.element.body
for el in list(body):
    if el.tag.endswith('}sectPr'):
        continue
    body.remove(el)

# ---- landscape Letter ----
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Emu(10058400), Emu(7772400)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Inches(0.75))
sec.top_margin = Inches(0.6)
sec.bottom_margin = Inches(0.6)
CONTENT_W = 11.0 - 1.5   # inches

# ---- footer tab stops: the template's right tab sits at the portrait text width; move it to landscape ----
ftr_p = sec.footer.paragraphs[0]
pPr = ftr_p._p.get_or_add_pPr()
tabs = pPr.find(qn('w:tabs'))
if tabs is None:
    tabs = OxmlElement('w:tabs'); pPr.append(tabs)
for pos in ("5040", "10080"):
    t_ = OxmlElement('w:tab'); t_.set(qn('w:val'), 'clear'); t_.set(qn('w:pos'), pos); tabs.append(t_)
t_ = OxmlElement('w:tab'); t_.set(qn('w:val'), 'right'); t_.set(qn('w:pos'), str(int(CONTENT_W * 1440))); tabs.append(t_)

# ---- helpers ----
def add_runs(par, text, size=None, color=None, bold=None):
    """Text with «citation» segments rendered small and gray."""
    parts = re.split(r'(«[^»]*»)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('«'):
            r = par.add_run(part[1:-1])
            r.font.size = Pt((size or 10) - 1.5)
            r.font.color.rgb = CITE_COLOR
        else:
            r = par.add_run(part)
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if bold is not None: r.font.bold = bold
    return par

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge, (val, sz, color) in edges.items():
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), val); el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        borders.append(el)

def set_cell_margins(cell, top=50, bottom=50, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f'w:{k}'); el.set(qn('w:w'), str(v)); el.set(qn('w:type'), 'dxa'); mar.append(el)
    tcPr.append(mar)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit'); trPr.append(el)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); trPr.append(el)

def make_table(headers, rows, widths_in, shade_cols=(), bold_last=False, font_pt=TABLE_PT):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = doc.styles['Normal Table']
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    # fixed layout
    tblPr = t._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(sum(widths_in) * 1440))); tblW.set(qn('w:type'), 'dxa')
    # header
    hdr = t.rows[0]
    repeat_header(hdr); no_split(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; c.width = Inches(widths_in[i])
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        add_runs(c.paragraphs[0], h, size=font_pt, bold=True)
        set_cell_borders(c, bottom=('single', 8, '404040'))
        set_cell_margins(c)
    # body
    for ri, row in enumerate(rows):
        r = t.add_row(); no_split(r)
        last = bold_last and ri == len(rows) - 1
        for i, txt in enumerate(row):
            c = r.cells[i]; c.width = Inches(widths_in[i])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, txt, size=font_pt, bold=(True if last else None))
            set_cell_borders(c, bottom=('single', 4, RULE_COLOR))
            set_cell_margins(c)
            if i in shade_cols:
                set_cell_shading(c, EFFECT_FILL)
    # column widths (grid + every cell)
    grid = t._tbl.find(qn('w:tblGrid'))
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths_in[i] * 1440)))
    for i, w in enumerate(widths_in):
        for cell in t.columns[i].cells:
            cell.width = Inches(w)
    return t

def heading(text):
    p = doc.add_paragraph(text, style='Heading 1')
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p

def body_par(text, size=10, after=6):
    p = doc.add_paragraph()
    add_runs(p, text, size=size)
    p.paragraph_format.space_after = Pt(after)
    return p

# ---- content ----
doc.add_paragraph("Strev Amendment No. 2 (AI Features): What Changed and What It Costs", style='Title')

body_par("A head start on what changed in the draft, not the review itself: the legal review is Brian’s to lead, and my read-through points to potential effects on the FY27 AOP and the Strev renewal BIR that should be reviewed but are not sized here.", size=10, after=3)
body_par("References are page and section in: the Draft (received 2026-10-21); the Agreement (White Label Software License and Service Agreement, 2025-03-21); Amend. 1 (2025-08-13); v2.5 (Revnue pricing framework, May 2026). Shaded columns carry the difference.", size=9, after=2)

p = doc.add_paragraph(style='Bottom Line')
add_runs(p, "Key takeaway: ", size=10.5, bold=True)
add_runs(p, "the draft defines the capability set as “AI Features” and prices it at $164,000 per Environment per year through two mandatory fees (AI Platform License, AI Feature Pack Base License); every internal number to date has used $84,000. It also changes the agreement’s payment terms for these fees and charges every Environment, including QA, where the agreement treats non-production use as included.", size=10.5)
p.paragraph_format.space_after = Pt(6)

heading("1. Financial changes")
fin_headers = ["#", "Change", "Before", "Draft Amendment No. 2", "Effect"]
fin_rows = [
 ["1", "Second mandatory fee",
  "One $84,000 base license per environment «(v2.5 p. 13, §6.1)»; §6 states the earlier platform-fee model was retired.",
  "$80,000 AI Platform License plus $84,000 AI Feature Pack Base License; both must be paid «(Draft p. 4, Exhibit A table and term 1)».",
  "+$80,000 per Environment; a return to a fee layer the vendor said it had dropped."],
 ["2", "AI usage paid by Wesco and its customers",
  "Revnue credits, $10,000 minimum per tenant; Revnue-hosted local model at no token cost «(v2.5 pp. 9-10, §5.1)»; usage statements per tenant «(p. 15, §6.7)».",
  "No Revnue usage charge. Wesco must hold and pay its own OpenAI, Anthropic or Google account; no local model; no usage reporting «(Draft p. 2, §2.12(b))».",
  "Uncapped cost outside the contract; one Wesco account serves every tenant and End User, with no way to charge it back."],
 ["3", "Support uplift",
  "AI fees not addressed (v2.5 silent).",
  "AI Fees count toward Annual Contract Value «(Draft p. 4, term 5)»; Standard support is 20% of ACV, Premium 30% «(Amend. 1 p. 4)».",
  "+$32,800 (Standard) or +$49,200 (Premium) per Environment."],
 ["4", "Charged per Environment, QA included",
  "Non-production use is part of the license; extra environments at direct cost only, outside standard fees «(Agreement p. 3, §2.2(c))».",
  "An Environment is any multi-tenant deployment, Provider- or Customer-hosted; each needs both fees «(Draft p. 4, term 1)». No carve-out for test or QA.",
  "A second $164,000 plus support if AI is installed in QA; conflicts with §2.2(c)."],
 ["5", "Payment terms changed",
  "Fees invoiced monthly in arrears, net 30 after an undisputed invoice «(Agreement p. 6, §4.1)»; half of subscription fees drawn from the $500,000 prepayment «(p. 16, Exhibit A §1.1)».",
  "AI Fees invoiced annually in advance, net 30 from receipt; no deduction from the prepayment «(Draft p. 4, terms 2 and 3)».",
  "A change to the agreement’s payment mechanics, not an entitlement change: a full year of cash at order; prepayment unusable."],
 ["6", "Renewal price increases",
  "Exhibit A prices fixed for the five-year term «(Amend. 1 pp. 3-4)»; no escalator in the agreement.",
  "Revnue may revise the AI Fees for any renewal on 90 days’ notice «(Draft p. 4, term 6)».",
  "First uncapped increase right in the relationship, from year two."],
]
make_table(fin_headers, fin_rows, [0.3, 1.3, 2.75, 2.85, 2.3], shade_cols=(4,))

h2 = heading("2. Total financial impact")
h2.paragraph_format.page_break_before = True   # Darryl, 2026-10-21: section 2 starts a new page
tot_headers = ["Item", "Previously assumed", "Draft", "Additional"]
tot_rows = [
 ["AI Features license, one Environment", "$84,000", "$164,000", "+$80,000"],
 ["Support on the AI Fees", "Not modeled", "$32,800 to $49,200", "+$32,800 to $49,200"],
 ["One Environment, year one", "$84,000", "$196,800 to $213,200", "+$112,800 to $129,200"],
 ["QA Environment, if AI is installed there", "Direct hosting cost only", "$196,800 to $213,200", "+$196,800 to $213,200"],
 ["AI usage (OpenAI, Anthropic or Google)", "Inside v2.5 credits, $10,000 per tenant minimum", "Paid directly by Wesco; unsized", "Unknown"],
 ["Both Environments, year one", "$84,000", "$393,600 to $426,400 plus usage", "+$309,600 to $342,400 plus usage"],
]
make_table(tot_headers, tot_rows, [3.2, 2.1, 2.1, 2.1], shade_cols=(3,), bold_last=True)
body_par("Ranges reflect Standard (20%) or Premium (30%) support. Year-one amounts are payable in advance at order, outside the prepayment, and non-cancellable for the year.", size=9.5, after=2)

heading("3. Contract-term changes")
term_headers = ["#", "Term", "Before", "Draft Amendment No. 2", "Effect"]
term_rows = [
 ["7", "Scope: ten capabilities",
  "v2.5 base pack of five «(pp. 9-10, §5.1)»; Preventative Maintenance and IoT were separate packs «(p. 11)»; the chat assistant and LLM data upload were private-beta add-ons «(p. 12, §5.4)».",
  "(a) Anomaly Detection; (b) Preventative Maintenance Scheduling; (c) IoT Features; (d) Document Intelligence; (e) AI-Based Asset Filtering; (f) Document Extraction to Contracts, Assets and Licenses; (g) Content Filtering; (h) AI Enrichment; (i) Strev AI natural-language assistant; (j) AI Bulk Import and Update «(Draft p. 1, §1.19)».",
  "More in the bundle, which is part of the price. Downside: (c), (e) and (h) are undefined; the v2.5 guardrails and local model are not listed; the Software Service line reads “AI Features [(in development)]” «(Draft p. 1, §2.1)»."],
 ["8", "The features we asked for",
  "Data extraction, data normalization and the Virtual Site Survey (Amy’s three).",
  "Data extraction maps to (f). Data normalization is not named; (h) may be it. The Virtual Site Survey is not in the draft at all.",
  "Two of three unconfirmed; the site survey stays on its own pricing track."],
 ["9", "Minimum term and exit",
  "Convenience termination on 90 days’ notice «(Agreement p. 8, §9.2(d))»; term ends 2030-03-21 «(§9.1)».",
  "One-year non-cancellable, non-refundable minimum per Environment, then annual auto-renewal on 30 days’ notice «(Draft p. 4, term 4)».",
  "Convenience exit forfeits the year; nothing ties the AI term to the agreement’s end date."],
 ["10", "Liability on the AI leg",
  "Provider warranty «(Agreement p. 6, §5.2)» and indemnity including data breach «(p. 9, §10.1)» cover the Services.",
  "AI provider is Wesco’s vendor, not a Subcontractor; no Revnue warranty or indemnity for its outages, data handling or output; no accuracy warranty on AI Output «(Draft p. 2, §2.12(f), (g))».",
  "Breach, data-protection and accuracy exposure sits with Wesco, and with its End User terms once AI is resold."],
 ["11", "Kept",
  "n/a",
  "No training on Customer Data; AI Output is Customer Data; AI source code into escrow «(Draft pp. 2-3, §2.12(d), (e), (h))».",
  "Acceptable as drafted."],
]
make_table(term_headers, term_rows, [0.3, 1.3, 2.75, 2.85, 2.3], shade_cols=(4,))


doc.save(OUT)
print("saved", OUT)
