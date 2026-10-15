# v2 2026-10-14: CDA slide updated after the Matt Powers meeting (TS backing, Intel as potential anchor).
# Sean McCloskey portfolio deck, week of 2026-10-19 (python-pptx on the Wesco template).
# Run from a folder containing template.pptx (copy of Reference/Templates/Wesco_Slide_Template.pptx).
# Source of truth for wording: this file. Fold PowerPoint hand edits back in here.
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=C(0x1D,0x25,0x2D); GREEN=C(0x00,0x95,0x3A); SLATE=C(0x5B,0x7A,0x8E)
MIST=C(0xDB,0xE3,0xE9); GREY=C(0x6E,0x73,0x78); WHITE=C(0xFF,0xFF,0xFF)
INK=C(0x2B,0x2B,0x2B); SOFT=C(0xC9,0xD1,0xD6); PALE=C(0xE9,0xF3,0xEC)
MBLUE=C(0x00,0x49,0x86); MBLUE_TINT=C(0xCC,0xDA,0xE7); STEEL=C(0xF1,0xF5,0xF7); STEEL_MID=C(0xDF,0xE8,0xED)
AMBER=C(0xC9,0x7A,0x00)
F='Arial'
prs=Presentation('template.pptx')
ids=prs.slides._sldIdLst
for sid in list(ids):
    prs.part.drop_rel(sid.get(qn('r:id'))); ids.remove(sid)
L={l.name:l for l in prs.slide_layouts}

def add(layout,t=None):
    s=prs.slides.add_slide(L[layout])
    if t: s.shapes.title.text_frame.paragraphs[0].add_run().text=t
    return s
def tb(s,x,y,w,h):
    tf=s.shapes.add_textbox(In(x),In(y),In(w),In(h)).text_frame; tf.word_wrap=True
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,0)
    return tf
def run(tf,text,size,bold=False,color=INK,after=2,first=False,italic=False,align=PP_ALIGN.LEFT):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after=Pt(after); p.alignment=align
    r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=F; f.italic=italic
    return p
def box(s,x,y,w,h,fill=None,line=None,lw=1.25,adj=0.08,dash=False,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh=s.shapes.add_shape(shape,In(x),In(y),In(w),In(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb=line; sh.line.width=Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE as D; sh.line.dash_style=D.DASH
    sh.shadow.inherit=False
    return sh
def label(sh,lines,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,m=0.12):
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=In(m); tf.margin_top=tf.margin_bottom=In(0.06)
    for i,(t,sz,b,c) in enumerate(lines): run(tf,t,sz,b,c,after=2,first=(i==0),align=align)
    return sh
def eyebrow(s,x,y,w,text,color=GREEN):
    run(tb(s,x,y,w,0.25),text.upper(),10,True,color,after=0,first=True)
def conn(s,x1,y1,x2,y2,color=SLATE,lw=1.75,arrow=True):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(x1),In(y1),In(x2),In(y2))
    c.line.color.rgb=color; c.line.width=Pt(lw)
    if arrow:
        ln=c.line._get_or_add_ln(); te=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}); ln.append(te)
    return c
def takeaway(s,text,y=5.95,h=0.5):
    b=box(s,0.73,y,11.86,h,PALE,adj=0.12)
    label(b,[(text,13,True,INK)],align=PP_ALIGN.LEFT,m=0.22)
def note(s,t): s.notes_slide.notes_text_frame.text=t
def table(s,x,y,w,data,col_w,row_h,size=12,head=12,head_fill=NAVY,first_bold=True,center=False):
    gf=s.shapes.add_table(len(data),len(data[0]),In(x),In(y),In(w),In(row_h*len(data))); t=gf.table
    for i,cw in enumerate(col_w): t.columns[i].width=In(cw)
    for ri,row in enumerate(data):
        t.rows[ri].height=In(row_h)
        for ci,val in enumerate(row):
            cell=t.cell(ri,ci); cell.margin_left=cell.margin_right=In(0.12)
            cell.margin_top=cell.margin_bottom=In(0.05); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=head_fill if ri==0 else (WHITE if ri%2 else STEEL)
            tf=cell.text_frame; tf.word_wrap=True
            for k,line in enumerate(str(val).split('\n')):
                p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
                r=p.add_run(); r.text=line; f=r.font; f.name=F
                f.size=Pt(head if ri==0 else size); f.bold=(ri==0 or (ci==0 and first_bold))
                f.color.rgb=WHITE if ri==0 else (NAVY if ci==0 else INK)
                if center and ci: p.alignment=PP_ALIGN.CENTER
    return t

# 1 ------------------------------------------------------------ title
s=add('Title Slide')
for ph in list(s.placeholders):
    i=ph.placeholder_format.idx
    if i==12: ph.text_frame.paragraphs[0].add_run().text='CSS Platforms'
    elif i==1: ph.text_frame.paragraphs[0].add_run().text='October 2026'
    elif i not in (13,14) and ph.has_text_frame and not ph.text_frame.text.strip():
        ph._element.getparent().remove(ph._element)
note(s,"30 minutes. WAM first, about 15 minutes, because it carries the only decisions. Then entroCIM's AI and the Ascent position, about 10, and CDA, about 5. The written note sent ahead carries the three asks.")

# 2 ------------------------------------------------------------ portfolio
s=add('Blank-with-Title','Portfolio')
data=[['Product','Where it stands','Next'],
 ['entroCIM','Reduced capacity from 2026 to 2027.\nAbsorbing CDA.','2027 roadmap reconciled with CDA.\nWoolly AI and Ascent as the commercial levers.'],
 ['WAM','Production capability holds at any staffing level.\nCapacity reset in progress.','Roadmap reconciliation on scalable, valuable features.\nCapacity level agreed.'],
 ['Ascent Navigator','Airwall running at Ascent’s office.\nDevice data being gathered.','Show the value to Ascent’s team, then position entroCIM as a single pane of glass across its data center contracts.'],
 ['CDA','Built from the entroCIM team, no FY27 line.\nBill Muczko sponsoring.','Business review of the services with TS.']]
table(s,0.73,1.55,11.86,data,[2.3,4.6,4.96],0.86,size=13,head=13)
note(s,"One row per product, on the question that matters now. Only WAM needs a decision today. entroCIM, Ascent and CDA are positioning and awareness.\n\nentroCIM: capacity is reduced from this year to next. Ron, David and I are reconciling the 2027 core roadmap with building CDA, complete by the end of October.")

# 3 ------------------------------------------------------------ WAM today
s=add('Blank-with-Title','WAM today')
eyebrow(s,0.73,1.45,7.2,'Three work streams, set by the integration tools')
BW,BH,Y=1.85,0.95,1.95
xs=[0.73,3.58,6.43]
names=[('DDP / EUDM','Wesco enterprise data'),('ADM','Asset data product'),('Strev','Customer-facing platform')]
fills=[MBLUE,NAVY,MBLUE]
for x,(n,sub),fl in zip(xs,names,fills):
    label(box(s,x,Y,BW,BH,fl),[(n,16,True,WHITE),(sub,11,False,SOFT)])
for x1,x2,t in ((xs[0]+BW,xs[1],'Azure Databricks'),(xs[1]+BW,xs[2],'APIM')):
    conn(s,x1+0.04,Y+BH/2,x2-0.07,Y+BH/2)
    tf=tb(s,x1,Y+BH/2-0.52,x2-x1,0.45); tf.vertical_anchor=MSO_ANCHOR.BOTTOM
    for k,part in enumerate(t.split(' ')): run(tf,part,9.5,True,SLATE,after=0,first=(k==0),align=PP_ALIGN.CENTER)
streams=[('DDP integration','Brings orders and license data from Wesco systems into ADM.'),
         ('Core ADM product','License tracking, renewal notifications and data enrichment.'),
         ('Strev integration','License sync runs today. Asset sync needs the data in Wesco systems first.')]
for x,(h,b) in zip(xs,streams):
    label(box(s,x,3.12,BW,1.55,STEEL),[(h,12,True,NAVY),(b,11,False,INK)],align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP)
# scorecard
eyebrow(s,8.55,1.45,4.0,'Sales Operations scorecard, October')
stats=[('35','suppliers set up for license tracking'),('6,927','license parts in ADM'),('1,337','renewal notifications sent'),('238','sales reps reached')]
for i,(v,l) in enumerate(stats):
    x=8.55+(i%2)*2.1; y=1.95+(i//2)*1.38
    tf=tb(s,x,y,1.95,1.25); run(tf,v,30,True,GREEN,after=0,first=True); run(tf,l,11,False,INK,after=0)
run(tb(s,8.55,4.72,4.0,0.3),'North America. Reported by Sales Operations, not the product team.',9.5,False,GREY,first=True,italic=True)
takeaway(s,'ADM can only automate data Wesco systems already hold.',y=5.55)
note(s,"Work streams: DDP integration (ADM to EUDM on Azure Databricks), core ADM product development, and Strev integration (ADM to Strev on APIM). The integration tools define the skills the team needs.\n\nScorecard: Natalie Marquardt's October license-tracking update, Sales Operations. Coverage and use of a shipped capability: 35 suppliers (18 security, 17 IoT and other), 6,927 license parts, 1,337 notification line items to 238 North American reps. It is the one ADM capability with a business-side consumer reporting use every month, which is why notifications stay in every capacity option.\n\nIf asked about sales: the deck also shows about $697K in license sales for tracked suppliers, North America, August and September, 88% Milestone. It is not yet tied to the notifications, because nothing in DDP marks an order as a renewal. Sales Operations marks renewals by hand today.\n\nAsset sync: many asset programs run on Excel files built by warehouse and supply chain teams, uploaded to Strev by Managed Services. Near term the team helps automate those uploads. Long term the fix is capturing serial and MAC data in Wesco systems; Alan Martin's cross-SBU GPO session met on 10-15.")

# 4 ------------------------------------------------------------ WAM capacity
s=add('Blank-with-Title','WAM capacity')
data=[['','Run','Run plus DDP','Run plus completion'],
 ['License sync, notifications, production support','Kept','Kept','Kept'],
 ['Capacity added','None','Databricks data engineers to the level DDP needs','Plus backend, front end and APIM integration'],
 ['Enterprise data (DDP/EUDM)','Held','Advances','Advances'],
 ['Product and Strev completion','Extended substantially','Extended substantially','Planned pace']]
t=table(s,0.73,1.55,11.86,data,[3.5,2.55,3.06,2.75],0.68,size=12.5,head=13,center=True)
hc=t.cell(0,2); hc.fill.solid(); hc.fill.fore_color.rgb=GREEN
for ri in range(1,5):
    c=t.cell(ri,2); c.fill.solid(); c.fill.fore_color.rgb=PALE
    for p in c.text_frame.paragraphs:
        for r in p.runs: r.font.bold=True
eyebrow(s,0.73,5.12,8,'Skills the integration tools call for',SLATE)
chips=['Azure Databricks data engineers','APIM integration','Backend','Front end','QA']
x=0.73
for c_ in chips:
    w=0.32+0.085*len(c_)
    label(box(s,x,5.42,w,0.42,MBLUE_TINT,adj=0.3),[(c_,11,True,NAVY)],m=0.08); x+=w+0.15
takeaway(s,'Recommend Run plus DDP. Staffing numbers follow the WAM roadmap reconciliation.',y=6.0,h=0.46)
note(s,"Production capability, including renewal notifications, is kept at every level. The choice is how fast the enterprise data foundation moves.\n\nRecommendation: Run plus DDP. DDP engineers build the path Wesco-owned data takes into ADM, so it is the one stream where added people turn directly into that foundation. The data itself is owned outside the team, which is why the process changes matter as much as staffing. The Q1 2027 historical-data target comes off the plan; no committed date at any level.\n\nNo headcount on the slide on purpose: the stream numbers are being reconciled with Akansh and David in next week's roadmap reconciliation, which also tests each feature for scale and value.\n\nAsks for Sean:\n1. Where Digital Products taking over WAM development stands after his conversation with Akash. It decides whether this reset is our plan or our input to Merwan's team.\n2. Agreement in principle on Run plus DDP.\n3. His support with Dan Totino's team on flagging renewal orders when the rep creates them.")

# 5 ------------------------------------------------------------ entroCIM AI
s=add('Blank-with-Title','entroCIM AI')
frame=box(s,0.73,1.45,7.1,2.75,STEEL,adj=0.05)
run(tb(s,0.98,1.62,5,0.25),'Woolly inside entroCIM',11,True,SLATE,after=0,first=True)
cw,ch,cy=2.0,1.45,2.1
cxs=[0.98,3.28,5.58]
comps=[('Chat interface','Ask in plain language and review the answer'),('AI harness','Connects the model with skills and tools'),('Skills','Built-in platform knowledge plus custom skills')]
for x,(h,b) in zip(cxs,comps):
    label(box(s,x,cy,cw,ch,NAVY if h=='AI harness' else WHITE,line=None if h=='AI harness' else SOFT),
          [(h,14,True,WHITE if h=='AI harness' else NAVY),(b,10.5,False,SOFT if h=='AI harness' else INK)])
conn(s,cxs[0]+cw+0.03,cy+ch/2,cxs[1]-0.07,cy+ch/2)
conn(s,cxs[2]-0.03,cy+ch/2,cxs[1]+cw+0.07,cy+ch/2)
mx=cxs[1]; my=4.65
label(box(s,mx,my,cw,0.85,MBLUE),[('AI model provider',13,True,WHITE),('Model connection',10,False,SOFT)])
conn(s,mx+cw/2,cy+ch+0.03,mx+cw/2,my-0.07)
tf=tb(s,0.98,4.66,2.15,0.9); run(tf,'Bring your own keys',12,True,NAVY,first=True); run(tf,'The customer’s own API keys connect the model.',10.5,False,INK)
vals=[('Customers','More independence','Analyze their own data and build KPIs and dashboards with less SME help.'),
      ('Deployment','Less manual work','Help with setup, reports and issue investigation; skills reused across sites.'),
      ('Development','Faster development','Draft and refine code and views with less time on platform syntax.')]
for i,(e,h,b) in enumerate(vals):
    y=1.45+i*1.4
    bx=box(s,8.15,y,4.44,1.25,WHITE,line=SOFT)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=In(0.18); tf.margin_top=tf.margin_bottom=In(0.06)
    run(tf,e.upper(),9.5,True,GREEN,after=1,first=True); run(tf,h,13,True,NAVY,after=1); run(tf,b,10.5,False,INK,after=0)
run(tb(s,0.73,5.85,11.86,0.3),'Teams keep responsibility for validation. Adoption and time savings are expected benefits, not yet measured.',10,False,GREY,first=True,italic=True)
note(s,"Woolly is SkySpark's agent harness, shipped as a beta in SkySpark 4.0.6 and being built into entroCIM as a product capability. Three parts: the chat interface (conversation), the AI harness (coordination), and skills (expertise: built-in Haystack, Axon and platform knowledge plus custom skills the team writes). The harness calls an AI model provider using the customer's own API keys.\n\nValue, from Ron Snyder's framing: customers do more on their own, deployment teams spend less time on routine platform work, developers get a faster first draft. Expected, not measured.\n\nApplication opportunity: connect new or existing apps to AI workflows and chatbots powered by custom agents.\n\nIf asked: Woolly does not support Azure AI Foundry today, so development and test use direct API keys with their own small budget line. Evaluation runs with no production or customer data.")

# 6 ------------------------------------------------------------ Ascent
s=add('Blank-with-Title','Ascent')
steps=[('1','Connect','Airwall on site and running','Done',GREEN,WHITE),
       ('2','Collect','PDU, camera, BMS and metering data as available','In progress',MBLUE,WHITE),
       ('3','Show value','Graphics, KPIs, alarming and FDD in entroCIM','Next',STEEL_MID,NAVY),
       ('4','Position','Ascent’s team positions entroCIM as a single pane of glass across its data center contracts','Goal',WHITE,NAVY)]
SW,SH,SY=2.72,2.6,1.75
for i,(n,h,b,st,fl,tc) in enumerate(steps):
    x=0.73+i*(SW+0.32)
    bx=box(s,x,SY,SW,SH,STEEL if i<3 else PALE,adj=0.05)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.2); tf.margin_top=In(0.22)
    run(tf,n,28,True,GREEN if i==3 else MBLUE,after=4,first=True); run(tf,h,16,True,NAVY,after=4); run(tf,b,12,False,INK,after=0)
    chipb=box(s,x+0.2,SY+SH-0.58,1.35,0.36,fl,line=GREEN if st=='Goal' else None,adj=0.3)
    label(chipb,[(st,10.5,True,tc)],m=0.05)
    if i<3: conn(s,x+SW+0.03,SY+SH/2,x+SW+0.32-0.07,SY+SH/2,color=GREEN)
takeaway(s,'Prove the value at Ascent’s own office, then carry it to every Ascent data center contract.',y=4.85)
note(s,"The Airwall at Ascent's office is on site and running. The team is gathering PDU, camera, BMS and metering device information as it becomes available, then deploying graphics, KPIs, alarming and fault detection and diagnostics (FDD) in entroCIM to show the value to the Ascent team. The goal is for Ascent to position entroCIM as a single pane of glass across all of its data center contracts.\n\nThis is the interim step ahead of the Navigator integration on DJ Vaughn's roadmap (entroCIM integration support Q4 2027). It builds comfort with Ascent Ops now, inside FY27, so the later integration improves on how Ascent actually operates.\n\nCameras in scope also give CDA a working example inside a data center operation.")

# 7 ------------------------------------------------------------ CDA
s=add('Blank-with-Title','CDA')
cards=[('Position','A managed service on the devices Wesco already sells. One annual fee per site, vendor-neutral, cameras first. It stands on its own in the market.'),
       ('Build','Built by the entroCIM team on the entroCIM roadmap. No FY27 line.'),
       ('Sponsor and support','Bill Muczko sponsoring. Matt Powers (VP, Technology Services) backs the service and an operations-center delivery model.'),
       ('Anchor customers','Kroger and Cisco, from Bob Caprise; the plan is built on them. Intel is a potential anchor: TS writes its smart building spec, and Matt sees it as the kind of company that would love this service.')]
CW,CH=3.75,1.95
for i,(h,b) in enumerate(cards):
    x=0.73+(i%2)*(CW+0.3); y=1.45+(i//2)*(CH+0.3)
    bx=box(s,x,y,CW,CH,STEEL,adj=0.05)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.2); tf.margin_top=In(0.2)
    run(tf,h.upper(),10,True,GREEN,after=4,first=True); run(tf,b,12,False,INK,after=0)
nx=0.73+2*CW+0.6; nw=12.59-nx
nb=box(s,nx,1.45,nw,2*CH+0.3,NAVY,adj=0.04)
tf=nb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
tf.margin_left=tf.margin_right=In(0.25); tf.margin_top=In(0.25)
run(tf,'NEXT',10,True,C(0x6F,0xD0,0x8C),after=8,first=True)
for i,t_ in enumerate(['Business review of the services with TS, including replacing failed devices through the account\u2019s integrator','Right-size the entroCIM roadmap to build CDA with the existing team','Enterprise pricing for large multi-site rollouts']):
    run(tf,f'{i+1}   {t_}',14,True,WHITE,after=14)
note(s,"Summary only; not a decision for today.\n\nCDA was not approved as a standalone FY27 investment. Development continues from the entroCIM team, on the entroCIM roadmap. Market positioning is unchanged: CDA stands on its own and does not read as an entroCIM add-on.\n\nTS (formerly TSS): Emily Munk was supportive and pointed me to Matt Powers, VP of Technology Services. Matt and I met 10-14. He called CDA a great way to get into accounts, with pull-through to ALM and WAM, and the kind of sticky services model TS needs as hardware margins compress. He proposed running it from an operations center, the same model I had in mind: a few people in North America, a couple in Europe, most of the work in India, with in-market people for larger remediation. It fits his push to move TS up the stack. Next is a business review of the services with Matt and Cody Remington, including how a failed device is replaced and dispatched through the integrator on the account. That settles who delivers and supports CDA.\n\nIntel: TS writes Intel's smart building specification, and Matt said a company like Intel would love this service. That relationship makes it a potential anchor alongside Kroger and Cisco.\n\nAnchor customers came from Bob Caprise and are how the plan was built. Peter Ferrell's challenge on enterprise pricing for large multi-site rollouts (banks, thousands of small branches) is still being worked.\n\nRon, David and I are reconciling the 2027 core entroCIM roadmap with building CDA, complete by the end of October.")

# 8 ------------------------------------------------------------ backup: ADM constraints
s=add('Blank-with-Title','ADM constraints')
eyebrow(s,0.73,1.3,4,'Backup',SLATE)
data=[['Constraint','Use case and impact','Ask','Owner','Status'],
 ['Trimmed part numbers in EUDM','License orders from EUDM not pulled into ADM\nImpact: lost Sales Ops data and renewal opportunities','Upstream EUDM enhancement','EUDM/AMS','In analysis'],
 ['Null WAM shipping and country data','Managed services data prep for Strev upload\nImpact: no warranty dates, so missed renewals','UDM/EUDM root cause fix','UDM','In analysis'],
 ['Missing license/asset durations','Tracking products ADM is configured for\nImpact: incomplete customer view of Wesco data','Supplier data hydration (via Jarod)','Supplier relationship team','To coordinate'],
 ['Missing serial/MAC ID data','Asset tracking\nImpact: lost tracking and renewal opportunity','GPO alignment on to-be process','WMX/VAS GPOs','In analysis'],
 ['No parent-child relationships','Component view of complex, nested builds\nImpact: no full assembly visibility','VAS API or EUDM data feed','VAS','Not started'],
 ['Unidentifiable renewal orders','Identify renewing orders to count renewal dollars\nImpact: renewals cannot be marked accurately','Implement renewal flag','Dan Totino / DDP GPOs','Not started']]
t=table(s,0.73,1.65,11.86,data,[2.35,4.2,2.2,1.71,1.4],0.66,size=10.5,head=11)
SC={'In analysis':MBLUE,'To coordinate':AMBER,'Not started':GREY}
for ri in range(1,7):
    c=t.cell(ri,4); r=c.text_frame.paragraphs[0].runs[0]; r.font.color.rgb=SC[r.text]; r.font.bold=True
    for p in c.text_frame.paragraphs: p.alignment=PP_ALIGN.LEFT
    p=t.cell(ri,1).text_frame.paragraphs
    if len(p)>1:
        rr=p[1].runs[0]; txt_=rr.text; rr.text=txt_.split(':',1)[1]
        b=p[1].runs[0]._r.addprevious if False else None
        nr=p[1].add_run(); nr.text=rr.text; nr.font.name=F; nr.font.size=Pt(10.5); nr.font.color.rgb=INK
        rr.text='Impact:'; rr.font.bold=True; rr.font.color.rgb=NAVY
run(tb(s,0.73,6.33,11.86,0.25),'Source: ADM product team, October 2026. Every ask sits with an owner outside the WAM team.',9.5,False,GREY,first=True,italic=True)
note(s,"Only if Sean wants the detail behind 'upstream data.' All six asks sit with owners outside the WAM team; the team's mitigations are investigations and workarounds. Renewal-order flagging starts with Dan Totino. Serial and MAC capture is in Alan Martin's cross-SBU GPO effort.")

out='2026-10-14_Sean_Portfolio_Deck_v2.pptx'
prs.save(out); print('saved',out,len(prs.slides._sldIdLst))
