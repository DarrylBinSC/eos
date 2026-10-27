# CSS Platforms portfolio slide, 2026-10-22 (python-pptx on the Wesco template).
# v3 2026-10-22: CDA sponsor Bill Muczko (Darryl).
# v2 2026-10-22: card titled WAM / ADM (no subline); CAT and i3 footer removed (Darryl).
# One card per live product in Products/_products.json, same lines on each. Source of truth for wording: this file.
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
NAVY=C(0x1D,0x25,0x2D); GREEN=C(0x00,0x95,0x3A); WHITE=C(0xFF,0xFF,0xFF); INK=C(0x2B,0x2B,0x2B)
STEEL=C(0xF1,0xF5,0xF7); GREY=C(0x6E,0x73,0x78); AMBER=C(0xC9,0x7A,0x00); MBLUE=C(0x00,0x49,0x86); LGREEN=C(0x6F,0xD0,0x8C)
F='Arial'
prs=Presentation('template.pptx'); ids=prs.slides._sldIdLst
for sid in list(ids): prs.part.drop_rel(sid.get(qn('r:id'))); ids.remove(sid)
L={l.name:l for l in prs.slide_layouts}
def tb(s,x,y,w,h):
    tf=s.shapes.add_textbox(In(x),In(y),In(w),In(h)).text_frame; tf.word_wrap=True
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,0)
    return tf
def run(tf,text,size,bold=False,color=INK,after=2,first=False,italic=False,align=PP_ALIGN.LEFT):
    p=tf.paragraphs[0] if first else tf.add_paragraph(); p.space_after=Pt(after); p.alignment=align
    r=p.add_run(); r.text=text; f=r.font; f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=F; f.italic=italic; return p
def box(s,x,y,w,h,fill=None,line=None,lw=1.25,adj=0.08):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,In(x),In(y),In(w),In(h)); sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def label(sh,lines,align=PP_ALIGN.CENTER,m=0.1):
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=In(m); tf.margin_top=tf.margin_bottom=In(0.04)
    for i,(t,sz,b,c) in enumerate(lines): run(tf,t,sz,b,c,after=0,first=(i==0),align=align)
    return sh

s=prs.slides.add_slide(L['Blank-with-Title'])
s.shapes.title.text_frame.paragraphs[0].add_run().text='CSS Platforms'

# (name, sub, what, serves, stage, stage colour, 2027 headline, sponsor)
cards=[
 ('entroCIM','', 'Infrastructure lifecycle intelligence for facilities and data centers: a customer’s building and IT systems integrated into one platform for monitoring, analytics and action.',
  'Facility and data center operators; installed base includes Federal Reserve, Ford, Dell and Disney.',
  'In market',GREEN,'entroCIM 4 and entroCIM AI launch in Q1 2027.','Dirk Naylor'),
 ('WAM / ADM','', 'One governed record of what is installed: licenses, warranties, coverage and lifecycle, built from Wesco’s own data. ADM is the asset-data home within DDP.',
  'Sales, Managed Services and their customers.',
  'In market',GREEN,'Asset records for Managed Services customers in Q1 2027.','Brad Stevens'),
 ('Navigator','', 'Operations platform for mission-critical facilities: asset management, work orders, budgeting and forecasting, replacing Maximo.',
  'Ascent’s data center operations, as the service wrapper on CSS Platforms.',
  'In market',GREEN,'entroCIM feeds Navigator, removing manual meter entry; a single pane of glass across Ascent’s contracts.','Dirk Naylor'),
 ('CDA','', 'Connected Device Assurance: an annual per-site managed service that keeps a customer’s connected devices online, current and covered, run from a Wesco operations center.',
  'Owners of camera, access control, IoT and OT estates; delivered with Technology Services.',
  'In build',AMBER,'Built from the entroCIM team under the FY27 plan; first sites follow a delivery owner in TSS.','Bill Muczko'),
 ('Launch Studio','', 'Turns business-unit prototypes into supported digital products: an onboarding review, a build to standard, and a run model with retirement at end of life.',
  'Business units building with AI tools; Digital Products as the home for what passes review.',
  'In exploration',MBLUE,'Stand up the onboarding review and the first productized idea in 2027.','Sean McCloskey'),
]
X0=0.73; GAP=0.18; W=(11.86-GAP*4)/5; Y=1.45; H=4.5
for i,(n,sub,what,serves,stage,sc,head,spon) in enumerate(cards):
    x=X0+i*(W+GAP)
    box(s,x,Y,W,H,STEEL,adj=0.05)
    hb=box(s,x,Y,W,0.62,NAVY,adj=0.05)
    label(hb,[(n,15,True,WHITE)]+([(sub,9.5,True,LGREEN)] if sub else []),align=PP_ALIGN.LEFT,m=0.14)
    label(box(s,x+0.14,Y+0.74,1.25,0.26,sc,adj=0.4),[(stage,8.5,True,WHITE)],m=0.04)
    tf=tb(s,x+0.14,Y+1.12,W-0.28,H-1.3)
    def sec(e,t,first=False):
        run(tf,e,8,True,GREEN,after=1,first=first); run(tf,t,9.5,False,INK,after=7)
    sec('WHAT IT IS',what,True); sec('WHO IT SERVES',serves); sec('2027',head)
    run(tf,'Sponsor: '+spon,8.5,False,GREY,after=0,italic=True)
prs.save('../2026-10-22_CSS_Platforms_Portfolio_v3.pptx'); print('saved')
