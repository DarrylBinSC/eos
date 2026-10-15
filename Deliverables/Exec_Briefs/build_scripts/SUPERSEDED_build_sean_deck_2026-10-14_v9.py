# v9 2026-10-14: WAM roadmap phase contents from Nitin's Estimation for WAM.xlsx.
# v8 2026-10-14: WAM roadmap slide added after WAM capacity, from the Akansh/Nitin email thread (09-04 to 10-14).
# v7 2026-10-14: 'Production features keep running' removed from the WAM portfolio row (Darryl).
# v6 2026-10-14: Darryl's QA answers. 'Development capacity reduced from 2026 to 2027' for entroCIM and WAM (AOP outcome); DC app in Q1 2027; AI Council approved, C2P started; 3.1.12.14 released by DCOE.
# v5 2026-10-14: QA pass. entroCIM casing in the Ascent eyebrow; scorecard counts labeled as line items and region assumption removed; 'commercial levers' wording; CDA operations center no longer stated as TS-owned; integrator replacement moved out of the customer outcome.
# v4 2026-10-14: entroCIM roadmap slide added after entroCIM AI; Woolly AI launch Q1 2027 with entroCIM 4 (SkySpark 4).
# v3 2026-10-14: Darryl's slide-by-slide review (portfolio wording, WAM capacity rebuilt, 30+ skills, Ascent tied to the joint plan, CDA as a service flow).
# v2 2026-10-14: CDA slide updated after the Matt Powers meeting (TS backing, Intel as potential anchor).
# Sean McCloskey portfolio deck, week of 2026-10-19 (python-pptx on the Wesco template).
# Run from a folder containing template.pptx (copy of Reference/Templates/Wesco_Slide_Template.pptx).
# Source of truth for wording: this file. Fold PowerPoint hand edits back in here.
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=C(0x1D,0x25,0x2D); GREEN=C(0x00,0x95,0x3A); SLATE=C(0x5B,0x7A,0x8E)
MIST=C(0xDB,0xE3,0xE9); GREY=C(0x6E,0x73,0x78); WHITE=C(0xFF,0xFF,0xFF)
INK=C(0x2B,0x2B,0x2B); SOFT=C(0xC9,0xD1,0xD6); PALE=C(0xE9,0xF3,0xEC)
MBLUE=C(0x00,0x49,0x86); MBLUE_TINT=C(0xCC,0xDA,0xE7); STEEL=C(0xF1,0xF5,0xF7); STEEL_MID=C(0xDF,0xE8,0xED)
AMBER=C(0xC9,0x7A,0x00)
F='Arial'
prs=Presentation('template.pptx')
ids=prs.slides._sldIdLst
for sid in list(ids):
    prs.part.drop_rel(sid.get(qn('r:id'))); ids.remove(sid)
L={l.name:l for l in prs.slide_layouts}

def add(layout,t=None):
    s=prs.slides.add_slide(L[layout])
    if t: s.shapes.title.text_frame.paragraphs[0].add_run().text=t
    return s
def tb(s,x,y,w,h):
    tf=s.shapes.add_textbox(In(x),In(y),In(w),In(h)).text_frame; tf.word_wrap=True
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,0)
    return tf
def run(tf,text,size,bold=False,color=INK,after=2,first=False,italic=False,align=PP_ALIGN.LEFT):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after=Pt(after); p.alignment=align
    r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=F; f.italic=italic
    return p
def box(s,x,y,w,h,fill=None,line=None,lw=1.25,adj=0.08,dash=False,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh=s.shapes.add_shape(shape,In(x),In(y),In(w),In(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb=line; sh.line.width=Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE as D; sh.line.dash_style=D.DASH
    sh.shadow.inherit=False
    return sh
def label(sh,lines,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,m=0.12):
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=In(m); tf.margin_top=tf.margin_bottom=In(0.06)
    for i,(t,sz,b,c) in enumerate(lines): run(tf,t,sz,b,c,after=2,first=(i==0),align=align)
    return sh
def eyebrow(s,x,y,w,text,color=GREEN,upper=True):
    run(tb(s,x,y,w,0.25),text.upper() if upper else text,10,True,color,after=0,first=True)
def conn(s,x1,y1,x2,y2,color=SLATE,lw=1.75,arrow=True):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(x1),In(y1),In(x2),In(y2))
    c.line.color.rgb=color; c.line.width=Pt(lw)
    if arrow:
        ln=c.line._get_or_add_ln(); te=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}); ln.append(te)
    return c
def takeaway(s,text,y=5.95,h=0.5):
    b=box(s,0.73,y,11.86,h,PALE,adj=0.12)
    label(b,[(text,13,True,INK)],align=PP_ALIGN.LEFT,m=0.22)
def note(s,t): s.notes_slide.notes_text_frame.text=t
def table(s,x,y,w,data,col_w,row_h,size=12,head=12,head_fill=NAVY,first_bold=True,center=False):
    gf=s.shapes.add_table(len(data),len(data[0]),In(x),In(y),In(w),In(row_h*len(data))); t=gf.table
    for i,cw in enumerate(col_w): t.columns[i].width=In(cw)
    for ri,row in enumerate(data):
        t.rows[ri].height=In(row_h)
        for ci,val in enumerate(row):
            cell=t.cell(ri,ci); cell.margin_left=cell.margin_right=In(0.12)
            cell.margin_top=cell.margin_bottom=In(0.05); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=head_fill if ri==0 else (WHITE if ri%2 else STEEL)
            tf=cell.text_frame; tf.word_wrap=True
            for k,line in enumerate(str(val).split('\n')):
                p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
                r=p.add_run(); r.text=line; f=r.font; f.name=F
                f.size=Pt(head if ri==0 else size); f.bold=(ri==0 or (ci==0 and first_bold))
                f.color.rgb=WHITE if ri==0 else (NAVY if ci==0 else INK)
                if center and ci: p.alignment=PP_ALIGN.CENTER
    return t

# 1 ------------------------------------------------------------ title
s=add('Title Slide')
for ph in list(s.placeholders):
    i=ph.placeholder_format.idx
    if i==12: ph.text_frame.paragraphs[0].add_run().text='CSS Platforms'
    elif i==1: ph.text_frame.paragraphs[0].add_run().text='October 2026'
    elif i not in (13,14) and ph.has_text_frame and not ph.text_frame.text.strip():
        ph._element.getparent().remove(ph._element)
note(s,"30 minutes. WAM first, about 15 minutes, because it carries the only decisions. WAM is two slides plus the roadmap. Then entroCIM AI, its roadmap and the Ascent position, about 10, and CDA, about 5. The written note sent ahead carries the three asks.")

# 2 ------------------------------------------------------------ portfolio
s=add('Blank-with-Title','Portfolio')
data=[['Product','Where it stands','Next'],
 ['entroCIM','Development capacity reduced from 2026 to 2027.\nAbsorbing CDA.','2027 roadmap reconciled to new staffing and CDA.\nCommercial levers: Woolly AI (Q1 2027, with entroCIM 4) and Ascent.'],
 ['WAM','Development capacity reduced from 2026 to 2027.','Roadmap and skill mix right-sized to the smaller team.'],
 ['Ascent Navigator','Airwall running at Ascent’s office.\nDevice data being gathered.','Show the value to Ascent’s team, then position entroCIM as a single pane of glass across its data center contracts.'],
 ['CDA','Built from the entroCIM team, no FY27 line.\nBill Muczko sponsoring; Matt Powers (TS) supports.','Business review of the services with TS.']]
table(s,0.73,1.55,11.86,data,[2.3,4.6,4.96],0.86,size=13,head=13)
note(s,"One row per product, on the question that matters now. WAM carries the asks. entroCIM, Ascent and CDA are positioning and awareness.\n\nentroCIM: development capacity is reduced from 2026 to 2027 (AOP outcome). Ron, David and I are reconciling the 2027 core roadmap with building CDA, complete by the end of October. Woolly AI and the Ascent deployment are the two commercial levers.\n\nCDA: Matt Powers, VP of Technology Services, backs the service and the operations-center delivery model (meeting 10-14).")


# 3 ------------------------------------------------------------ WAM today
s=add('Blank-with-Title','WAM today')
eyebrow(s,0.73,1.45,7.2,'Three work streams, set by the integration tools')
BW,BH,Y=1.85,0.95,1.95
xs=[0.73,3.58,6.43]
names=[('DDP / EUDM','Wesco enterprise data'),('ADM','Asset data product'),('Strev','Customer-facing platform')]
fills=[MBLUE,NAVY,MBLUE]
for x,(n,sub),fl in zip(xs,names,fills):
    label(box(s,x,Y,BW,BH,fl),[(n,16,True,WHITE),(sub,11,False,SOFT)])
for x1,x2,t in ((xs[0]+BW,xs[1],'Azure Databricks'),(xs[1]+BW,xs[2],'APIM')):
    conn(s,x1+0.04,Y+BH/2,x2-0.07,Y+BH/2)
    tf=tb(s,x1,Y+BH/2-0.52,x2-x1,0.45); tf.vertical_anchor=MSO_ANCHOR.BOTTOM
    for k,part in enumerate(t.split(' ')): run(tf,part,9.5,True,SLATE,after=0,first=(k==0),align=PP_ALIGN.CENTER)
streams=[('DDP integration','Brings orders and license data from Wesco systems into ADM.'),
         ('Core ADM product','License tracking, renewal notifications and data enrichment.'),
         ('Strev integration','License sync runs today. Asset sync needs the data in Wesco systems first.')]
for x,(h,b) in zip(xs,streams):
    label(box(s,x,3.12,BW,1.55,STEEL),[(h,12,True,NAVY),(b,11,False,INK)],align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP)
# scorecard
eyebrow(s,8.55,1.45,4.0,'Sales Operations scorecard, October')
stats=[('35','suppliers set up for license tracking'),('6,927','license parts in ADM'),('1,337','renewal notification line items'),('238','sales reps reached')]
for i,(v,l) in enumerate(stats):
    x=8.55+(i%2)*2.1; y=1.95+(i//2)*1.38
    tf=tb(s,x,y,1.95,1.25); run(tf,v,30,True,GREEN,after=0,first=True); run(tf,l,11,False,INK,after=0)
run(tb(s,8.55,4.72,4.0,0.3),'Reported by Sales Operations, not the product team.',9.5,False,GREY,first=True,italic=True)
takeaway(s,'ADM can only automate data Wesco systems already hold.',y=5.55)
note(s,"Work streams: DDP integration (ADM to EUDM on Azure Databricks), core ADM product development, and Strev integration (ADM to Strev on APIM). The integration tools define the skills the team needs.\n\nScorecard: Natalie Marquardt's October license-tracking update, Sales Operations. Coverage and use of a shipped capability: 35 suppliers (18 security, 17 IoT and other), 6,927 license parts, 1,337 notification line items to 238 reps (region not stated on the scorecard; the sales figure is North America). It is the one ADM capability with a business-side consumer reporting use every month, which is why notifications stay in every capacity option.\n\nIf asked about sales: the deck also shows about $697K in license sales for tracked suppliers, North America, August and September, 88% Milestone. It is not yet tied to the notifications, because nothing in DDP marks an order as a renewal. Sales Operations marks renewals by hand today.\n\nAsset sync: many asset programs run on Excel files built by warehouse and supply chain teams, uploaded to Strev by Managed Services. Near term the team helps automate those uploads. Long term the fix is capturing serial and MAC data in Wesco systems; Alan Martin's cross-SBU GPO session meets 10-15; update this line with its outcome.")

# 4 ------------------------------------------------------------ WAM capacity
s=add('Blank-with-Title','WAM capacity')
cols=[('1','The team','Development capacity reduced from 2026 to 2027.','Whatever the final count, capacity is limited.'),
      ('2','The roadmap','Right-sized to what the smaller team can deliver.','Scalable, valuable features only. Production first, then the DDP data foundation, then product and Strev completion as capacity allows.'),
      ('3','The skills','Set by the integration tools.',None)]
CWID,CHT,CY=3.72,3.55,1.5
for i,(n,h,lead,body) in enumerate(cols):
    x=0.73+i*(CWID+0.35)
    bx=box(s,x,CY,CWID,CHT,STEEL if i<2 else WHITE,line=None if i<2 else SOFT,adj=0.04)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.25); tf.margin_top=In(0.25)
    run(tf,n,28,True,MBLUE,after=4,first=True); run(tf,h,17,True,NAVY,after=6); run(tf,lead,14,True,INK,after=6)
    if body: run(tf,body,12.5,False,INK,after=0)
    if i<2: conn(s,x+CWID+0.04,CY+CHT/2,x+CWID+0.35-0.07,CY+CHT/2,color=GREEN)
sx=0.73+2*(CWID+0.35)+0.25
for j,(c_,sub) in enumerate([('Azure Databricks engineers (EUDM)',''),('APIM integration (Strev)',''),('Backend',''),('Front end',''),('QA','')]):
    y=CY+1.68+j*0.36
    b=box(s,sx,y,CWID-0.5,0.33,MBLUE_TINT,adj=0.3)
    label(b,[((c_+('  ·  '+sub if sub else '')),10.5,True,NAVY)],align=PP_ALIGN.LEFT,m=0.12)
takeaway(s,'We will right-size the roadmap and the skill mix to the team we have. Staffing detail follows the WAM roadmap reconciliation.',y=5.4,h=0.62)
note(s,"Keep this simple. Development capacity is reduced from 2026 to 2027 (the AOP outcome), and whatever the final numbers, we will have limited resources.\n\nWhat we do about it: right-size the WAM roadmap to what the smaller team can deliver, keeping only features that scale and carry value. Production keeps running (license sync, renewal notifications, production support). After that, the DDP data foundation comes first, because ADM can only automate data Wesco systems hold. Product and Strev completion work extends.\n\nThe integration tools set the skill mix: Azure Databricks data engineers for the EUDM integration, APIM integration for Strev, plus backend, front end and QA. Akansh, David and I reconcile the roadmap and staffing next week.\n\nAsks for Sean:\n1. Where Digital Products taking over WAM development stands after his conversation with Akash.\n2. Agreement on right-sizing the roadmap to the smaller team, with the DDP foundation first after production.\n3. His support with Dan Totino's team on flagging renewal orders when the rep creates them.")


# 4b ----------------------------------------------------------- WAM roadmap
s=add('Blank-with-Title','WAM roadmap')
heads=[('Shipping now','Sales Operations and Managed Services',STEEL,NAVY,INK,MBLUE),
       ('Early 2027','Asset functionality for Managed Services',NAVY,WHITE,SOFT,C(0x6F,0xD0,0x8C)),
       ('Rest of 2027','Being reconciled to new staffing',STEEL,NAVY,INK,MBLUE)]
items=[['Renewal notifications to sales reps (live since August)','License sync from ADM to Strev (running)','WAM Data Enrichment for Strev uploads, 30 October','Orders beyond the configured supplier ID, 10 November'],
       ['January: asset records from Wesco orders, with warranty and coverage, serial numbers from WMX, and an asset grid with search and export','February: end-user verification, 30, 60 and 90-day expiry view, audit and edit','Asset upload by template lands ahead of both, in Q4','Dates are the delivery team’s estimate, held until the roadmap reconciliation'],
       ['DDP data foundation (EUDM integration on Azure Databricks)','Asset sync from ADM to Strev, once serial and MAC data is in Wesco systems','Renewal attribution, once orders carry a renewal flag','Historical order retrieval, no committed date']]
CW3,CH3,Y3=3.72,4.05,1.5
for i,((ey,h,fl,hc,bc,ec),its) in enumerate(zip(heads,items)):
    x=0.73+i*(CW3+0.35)
    bx=box(s,x,Y3,CW3,CH3,fl,adj=0.04)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.25); tf.margin_top=In(0.25)
    run(tf,ey.upper(),10.5,True,ec,after=3,first=True); run(tf,h,16,True,hc,after=10)
    for k,it in enumerate(its):
        last_note=(i==1 and k==len(its)-1)
        p=run(tf,it,(12 if i!=1 else 11.5) if not last_note else 10.5,(i==1 and not last_note),(bc if i==1 else INK) if not last_note else SOFT,after=6,italic=last_note)
        if last_note: continue
        pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(In(0.2)))); pPr.set('indent',str(-int(In(0.2))))
        bc_=pPr.makeelement(qn('a:buClr'),{}); bc_.append(bc_.makeelement(qn('a:srgbClr'),{'val':'00953A'})); pPr.append(bc_)
        pPr.append(pPr.makeelement(qn('a:buSzPct'),{'val':'100000'}))
        pPr.append(pPr.makeelement(qn('a:buFont'),{'typeface':'Arial'}))
        pPr.append(pPr.makeelement(qn('a:buChar'),{'char':'•'}))
    if i<2: conn(s,x+CW3+0.04,Y3+0.55,x+CW3+0.35-0.07,Y3+0.55,color=GREEN)
takeaway(s,'Sales Operations has the renewal signal today. Managed Services gets asset data in ADM early in 2027. Scale beyond that depends on upstream data.',y=5.75,h=0.5)
note(s,"Shipping now: the renewal notification suite has been live to reps since August (the scorecard on the WAM slide is its usage). License sync to Strev runs today. Two epics are on track per the delivery team (Nitin Pandey, 10-08): WAM Data Enrichment on 30 October, requested by Senthil and the Managed Services team to cut the turnaround on their Strev uploads, and Orders Beyond the Configured Supplier ID on 10 November.\n\nEarly 2027: the Asset Grid for Managed Services was held while Data Enrichment went first, then re-estimated. Nitin's phased plan (10-14, with his estimation spreadsheet): first production release 4 to 8 January with asset records pulled from Wesco orders in EUDM (sales order, purchase order, warranty and coverage), serial numbers from WMX, and the asset grid with filters, search (Wesco SO, part numbers, serial numbers), export and dashboard; final release 10 to 12 February with end-user verification and defaults, the 30, 60 and 90-day expiry view, audit and edit, status schedulers and full regression. Asset upload by template is a separate epic landing in Q4. The India team is furloughed 21 December to 3 January. Nitin flagged that priority or scope changes move the dates. I'm holding them as targets until next week's roadmap reconciliation confirms the team can staff them.\n\nNote: the spreadsheet still carries the historical data extraction epic through mid-February 2027; Akansh withdrew the Q1 2027 target on 10-14, so the slide shows no date.\n\nRest of 2027: reconciled next week on the scalable, valuable test. The DDP data foundation comes first after production. Asset sync to Strev and renewal attribution both wait on data owned outside the team (serial and MAC capture via the GPOs; a renewal flag on the order from Dan Totino's team). The historical-order retrieval target was withdrawn; no replacement date.")

# 5 ------------------------------------------------------------ entroCIM AI
s=add('Blank-with-Title','entroCIM AI')
frame=box(s,0.73,1.45,7.1,2.75,STEEL,adj=0.05)
run(tb(s,0.98,1.62,5,0.25),'Woolly inside entroCIM',11,True,SLATE,after=0,first=True)
cw,ch,cy=2.0,1.45,2.1
cxs=[0.98,3.28,5.58]
comps=[('Chat interface','Ask in plain language and review the answer'),('AI harness','Connects the model with skills and tools'),('Skills','30+ built-in skills plus custom skills the team writes')]
for x,(h,b) in zip(cxs,comps):
    label(box(s,x,cy,cw,ch,NAVY if h=='AI harness' else WHITE,line=None if h=='AI harness' else SOFT),
          [(h,14,True,WHITE if h=='AI harness' else NAVY),(b,10.5,False,SOFT if h=='AI harness' else INK)])
conn(s,cxs[0]+cw+0.03,cy+ch/2,cxs[1]-0.07,cy+ch/2)
conn(s,cxs[2]-0.03,cy+ch/2,cxs[1]+cw+0.07,cy+ch/2)
mx=cxs[1]; my=4.65
label(box(s,mx,my,cw,0.85,MBLUE),[('AI model provider',13,True,WHITE),('Model connection',10,False,SOFT)])
conn(s,mx+cw/2,cy+ch+0.03,mx+cw/2,my-0.07)
label(box(s,4.63,1.55,3.0,0.34,GREEN,adj=0.3),[('Launches Q1 2027 with entroCIM 4',10,True,WHITE)],m=0.06)
tf=tb(s,0.98,4.66,2.15,0.9); run(tf,'Bring your own keys',12,True,NAVY,first=True); run(tf,'The customer’s own API keys connect the model.',10.5,False,INK)
vals=[('Customers','More independence','Analyze their own data and build KPIs and dashboards with less SME help.'),
      ('Deployment','Less manual work','Help with setup, reports and issue investigation; skills reused across sites.'),
      ('Development','Faster development','Draft and refine code and views with less time on platform syntax.')]
for i,(e,h,b) in enumerate(vals):
    y=1.45+i*1.4
    bx=box(s,8.15,y,4.44,1.25,WHITE,line=SOFT)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=In(0.18); tf.margin_top=tf.margin_bottom=In(0.06)
    run(tf,e.upper(),9.5,True,GREEN,after=1,first=True); run(tf,h,13,True,NAVY,after=1); run(tf,b,10.5,False,INK,after=0)
run(tb(s,0.73,5.85,11.86,0.3),'Teams keep responsibility for validation. Adoption and time savings are expected benefits, not yet measured.',10,False,GREY,first=True,italic=True)
note(s,"Woolly is SkySpark's agent harness, shipped as a beta in SkySpark 4.0.6 and being built into entroCIM as a product capability. Launch target: Q1 2027, with entroCIM 4 (built on SkySpark 4). Three parts: the chat interface (conversation), the AI harness (coordination), and skills (expertise: 30+ built-in skills covering Haystack, Axon and platform knowledge, plus custom skills the team writes). The harness calls an AI model provider using the customer's own API keys.\n\nValue, from Ron Snyder's framing: customers do more on their own, deployment teams spend less time on routine platform work, developers get a faster first draft. Expected, not measured.\n\nApplication opportunity: connect new or existing apps to AI workflows and chatbots powered by custom agents.\n\nGovernance: AI Council has approved Woolly; Ron Snyder has started C2P. If asked: Woolly does not support Azure AI Foundry today, so development and test use direct API keys with their own small budget line.")

# 5b ----------------------------------------------------------- entroCIM roadmap
s=add('Blank-with-Title','entroCIM roadmap')
heads=[('Shipping now','Release 3.1.12.14',STEEL,NAVY,INK,MBLUE),
       ('Q1 2027','entroCIM 4',NAVY,WHITE,SOFT,C(0x6F,0xD0,0x8C)),
       ('Rest of 2027','Being reconciled to new staffing',STEEL,NAVY,INK,MBLUE)]
items=[['Batch Point Writer','Graphics Builder grouping','Tiles Gridboard text widget','Search (beta)','Platform and security improvements','Regular release updates to hosted customers'],
       ['Built on SkySpark 4','Containerized deployment','Data center app (DC UI redesign)','Woolly AI launch: 30+ built-in skills, customer’s own API keys'],
       ['HQ UI redesign go-live','Site Readiness app','Impact Analysis','CDA device and configuration assurance','Two major feature releases a year']]
CW3,CH3,Y3=3.72,4.05,1.5
for i,((ey,h,fl,hc,bc,ec),its) in enumerate(zip(heads,items)):
    x=0.73+i*(CW3+0.35)
    bx=box(s,x,Y3,CW3,CH3,fl,adj=0.04)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.25); tf.margin_top=In(0.25)
    run(tf,ey.upper(),10.5,True,ec,after=3,first=True); run(tf,h,17,True,hc,after=10)
    for it in its:
        p=run(tf,it,12.5,i==1,bc if i==1 else INK,after=6)
        pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(In(0.2)))); pPr.set('indent',str(-int(In(0.2))))
        bc_=pPr.makeelement(qn('a:buClr'),{}); bc_.append(bc_.makeelement(qn('a:srgbClr'),{'val':'00953A'})); pPr.append(bc_)
        pPr.append(pPr.makeelement(qn('a:buSzPct'),{'val':'100000'}))
        pPr.append(pPr.makeelement(qn('a:buFont'),{'typeface':'Arial'}))
        pPr.append(pPr.makeelement(qn('a:buChar'),{'char':'•'}))
    if i<2: conn(s,x+CW3+0.04,Y3+0.55,x+CW3+0.35-0.07,Y3+0.55,color=GREEN)
takeaway(s,'Woolly AI launches with entroCIM 4 in Q1 2027. The rest of 2027 is being reconciled to the new staffing level.',y=5.75,h=0.5)
note(s,"Shipping now: Release 3.1.12.14 has been released by DCOE and is rolling out. After approval: customer communication on day 1, hosted deployment on day 2, on-premises on day 5. Hosted customers now get release communications and regular updates as standard.\n\nQ1 2027: entroCIM 4, built on SkySpark 4, ships containerized (containerization is the required deployment model going forward), with the data center app (the DC UI redesign). Woolly AI launches with it: 30+ built-in skills, bring your own API keys. AI Council has approved; Ron has started C2P.\n\nRest of 2027: Ron, David and I are reconciling the 2027 roadmap to the new staffing level, including where CDA's device and configuration assurance features land. Agreed shape: two major feature releases (Q2 and Q4) with smaller security and improvement releases between; strategic integrations only. The HQ UI redesign go-live, Site Readiness and Impact Analysis moved from Q4 2026 into 2027 and are placed by the reconciliation (UI redesign timing to be confirmed with David and Ron).")

# 6 ------------------------------------------------------------ Ascent
s=add('Blank-with-Title','Ascent')
steps=[('1','Connect','Airwall on site and running','Done',GREEN,WHITE),
       ('2','Collect','PDU, camera, BMS and metering data as available','In progress',MBLUE,WHITE),
       ('3','Show value','Graphics, KPIs, alarming and FDD in entroCIM','Next',STEEL_MID,NAVY),
       ('4','Position','A single pane of glass across Ascent’s data center contracts','Goal',WHITE,NAVY)]
SW,SH,SY=2.72,1.95,1.45
for i,(n,h,b,st,fl,tc) in enumerate(steps):
    x=0.73+i*(SW+0.32)
    bx=box(s,x,SY,SW,SH,STEEL if i<3 else PALE,adj=0.05)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.2); tf.margin_top=In(0.15)
    run(tf,f'{n}  {h}',15,True,NAVY,after=3,first=True); run(tf,b,11.5,False,INK,after=0)
    label(box(s,x+0.2,SY+SH-0.5,1.3,0.32,fl,line=GREEN if st=='Goal' else None,adj=0.3),[(st,10,True,tc)],m=0.05)
    if i<3: conn(s,x+SW+0.03,SY+SH/2,x+SW+0.32-0.07,SY+SH/2,color=GREEN)
eyebrow(s,0.73,3.7,11,'PART OF THE JOINT NAVIGATOR AND entroCIM PLAN WITH ASCENT',SLATE,upper=False)
LY=4.55
ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(0.85),In(LY),In(12.45),In(LY)); ln.line.color.rgb=SOFT; ln.line.width=Pt(2.5)
band=box(s,0.85,LY-0.09,3.1,0.18,GREEN,adj=0.5)
miles=[(0.85,'Now through 2027','Interim link over the Airwall. Ascent Ops sees its equipment.',GREEN,True),
       (4.35,'June 2027','NavWork replaces Maximo.',MBLUE,False),
       (7.15,'Q4 2027','entroCIM integration with NavWork begins.',MBLUE,False),
       (9.95,'Q2 2028','Facility connectivity: meters and CM points.',MBLUE,False)]
for x,d,t,c,here in miles:
    if not here:
        dot=s.shapes.add_shape(MSO_SHAPE.OVAL,In(x-0.11),In(LY-0.11),In(0.22),In(0.22)); dot.fill.solid(); dot.fill.fore_color.rgb=c; dot.line.fill.background(); dot.shadow.inherit=False
    tf=tb(s,x if here else x-0.11,LY+0.25,2.6,0.9)
    run(tf,d,12.5,True,GREEN if here else NAVY,after=2,first=True); run(tf,t,11,False,INK,after=0)
run(tb(s,0.85,LY-0.42,3.0,0.25),'WE ARE HERE',9,True,GREEN,after=0,first=True)
takeaway(s,'Visibility for Ascent now; one connected portfolio with NavWork from Q4 2027.',y=5.75,h=0.5)
note(s,"The Airwall at Ascent's office is on site and running. The team is gathering PDU, camera, BMS and metering device information as available, then deploying graphics, KPIs, alarming and fault detection and diagnostics (FDD) to show the value to Ascent's team. The goal is for Ascent to position entroCIM as a single pane of glass across all of its data center contracts.\n\nThis is the interim step in the joint Navigator and entroCIM plan DJ Vaughn and I shared with Lisa Orton (September 2026): one link now through 2027 so Ascent Ops gains visibility inside this fiscal year; NavWork replaces Maximo by June 2027; entroCIM integration support with NavWork begins Q4 2027; facility connectivity (meters, CM points) Q2 2028. Dates are DJ's published Navigator roadmap. The integration that follows is designed against how Ascent actually operates.\n\nCameras in scope also give CDA a working example inside a data center operation.")

# 7 ------------------------------------------------------------ CDA
s=add('Blank-with-Title','CDA')
eyebrow(s,0.73,1.35,11,'A service powered by technology: Wesco runs it, the customer gets the outcome',SLATE)
NW,NH,NY=2.62,1.75,1.8
nodes=[('Customer site','Cameras, access control, IoT and OT devices Wesco already sold',WHITE,NAVY,INK,SOFT),
       ('CDA appliance','On site. Finds every device and connects securely back to Wesco',MBLUE,WHITE,SOFT,None),
       ('entroCIM in Wesco Azure','Inventory, firmware, cyber posture and alarms for every site',NAVY,WHITE,SOFT,None),
       ('Wesco operations center','Proposed home: TS. Mostly India, experts in North America and Europe. Fixes with customer approval',GREEN,WHITE,C(0xE6,0xF4,0xEA),None)]
for i,(h,b,fl,hc,bc,ln_) in enumerate(nodes):
    x=0.73+i*(NW+0.46)
    label(box(s,x,NY,NW,NH,fl,line=ln_),[(h,14,True,hc),(b,10.5,False,bc)],m=0.15)
    if i<3: conn(s,x+NW+0.04,NY+NH/2,x+NW+0.46-0.07,NY+NH/2,color=SLATE)
ob=box(s,0.73,3.85,11.86,0.62,PALE,adj=0.15)
label(ob,[('The customer gets a current inventory, patched firmware, secured devices, and problems fixed with their approval. One annual fee per site; nothing to buy or staff.',12,True,INK)],align=PP_ALIGN.LEFT,m=0.22)
pan=[('Backing','Bill Muczko sponsoring. Matt Powers (VP, Technology Services) backs the service and the operations-center model.'),
     ('Anchors','Kroger and Cisco, from Bob Caprise. Intel is a potential anchor: TS writes its smart building spec.'),
     ('Next','Business review with TS|Right-size the entroCIM roadmap to build CDA|Enterprise pricing for multi-site rollouts')]
PW=(11.86-0.6)/3
for i,(h,b) in enumerate(pan):
    x=0.73+i*(PW+0.3)
    bx=box(s,x,4.75,PW,1.55,STEEL if i<2 else NAVY,adj=0.05)
    tf=bx.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=tf.margin_right=In(0.2); tf.margin_top=In(0.15)
    run(tf,h.upper(),10,True,GREEN if i<2 else C(0x6F,0xD0,0x8C),after=4,first=True)
    for part in (b.split('|') if '|' in b else [b]): run(tf,part,11.5,False,INK if i<2 else WHITE,after=3)
note(s,"Summary only; not a decision for today.\n\nThe story left to right: Wesco has already sold the devices. A CDA appliance on site finds every device and connects securely back to entroCIM in Wesco Azure, where the operations team sees inventory, firmware, cyber posture and alarms across all its sites. A Wesco operations center runs it, mostly from India with experts in North America and Europe, and remediates with the customer's approval. TS is the proposed home; that is what the business review settles, so the slide does not state it as decided. The customer pays one annual fee per site and gets the outcome without buying a tool or adding staff. CDA stands on its own in the market; entroCIM runs underneath.\n\nCDA was not approved as a standalone FY27 investment. Development continues from the entroCIM team, on the entroCIM roadmap.\n\nMatt Powers, VP of Technology Services (10-14): a great way into accounts, pull-through to ALM and WAM, the sticky services model TS needs as hardware margins compress. He proposed the same operations-center model independently. Next is a business review with Matt and Cody Remington, including how a failed device is replaced and dispatched through the integrator on the account. That settles who delivers and supports CDA.\n\nAnchors came from Bob Caprise and are how the plan was built. Intel: TS writes its smart building specification, and Matt said a company like Intel would love this service, which makes it a potential anchor. Enterprise pricing for large multi-site rollouts (Peter Ferrell's challenge) is still being worked.")


# 8 ------------------------------------------------------------ backup: ADM constraints
s=add('Blank-with-Title','ADM constraints')
eyebrow(s,0.73,1.3,4,'Backup',SLATE)
data=[['Constraint','Use case and impact','Ask','Owner','Status'],
 ['Trimmed part numbers in EUDM','License orders from EUDM not pulled into ADM\nImpact: lost Sales Ops data and renewal opportunities','Upstream EUDM enhancement','EUDM/AMS','In analysis'],
 ['Null WAM shipping and country data','Managed services data prep for Strev upload\nImpact: no warranty dates, so missed renewals','UDM/EUDM root cause fix','UDM','In analysis'],
 ['Missing license/asset durations','Tracking products ADM is configured for\nImpact: incomplete customer view of Wesco data','Supplier data hydration (via Jarod)','Supplier relationship team','To coordinate'],
 ['Missing serial/MAC ID data','Asset tracking\nImpact: lost tracking and renewal opportunity','GPO alignment on to-be process','WMX/VAS GPOs','In analysis'],
 ['No parent-child relationships','Component view of complex, nested builds\nImpact: no full assembly visibility','VAS API or EUDM data feed','VAS','Not started'],
 ['Unidentifiable renewal orders','Identify renewing orders to count renewal dollars\nImpact: renewals cannot be marked accurately','Implement renewal flag','Dan Totino / DDP GPOs','Not started']]
t=table(s,0.73,1.65,11.86,data,[2.35,4.2,2.2,1.71,1.4],0.66,size=10.5,head=11)
SC={'In analysis':MBLUE,'To coordinate':AMBER,'Not started':GREY}
for ri in range(1,7):
    c=t.cell(ri,4); r=c.text_frame.paragraphs[0].runs[0]; r.font.color.rgb=SC[r.text]; r.font.bold=True
    for p in c.text_frame.paragraphs: p.alignment=PP_ALIGN.LEFT
    p=t.cell(ri,1).text_frame.paragraphs
    if len(p)>1:
        rr=p[1].runs[0]; txt_=rr.text; rr.text=txt_.split(':',1)[1]
        b=p[1].runs[0]._r.addprevious if False else None
        nr=p[1].add_run(); nr.text=rr.text; nr.font.name=F; nr.font.size=Pt(10.5); nr.font.color.rgb=INK
        rr.text='Impact:'; rr.font.bold=True; rr.font.color.rgb=NAVY
run(tb(s,0.73,6.33,11.86,0.25),'Source: ADM product team, October 2026. Every ask sits with an owner outside the WAM team.',9.5,False,GREY,first=True,italic=True)
note(s,"Only if Sean wants the detail behind 'upstream data.' All six asks sit with owners outside the WAM team; the team's mitigations are investigations and workarounds. Renewal-order flagging starts with Dan Totino. Serial and MAC capture is in Alan Martin's cross-SBU GPO effort.")

out='2026-10-14_Sean_Portfolio_Deck_v9.pptx'
prs.save(out); print('saved',out,len(prs.slides._sldIdLst))
